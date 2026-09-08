from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\explorax\Documents\Caza Estelar\Nueva carpeta")
SOURCE = ROOT / "INTEGRACION_ADMIN_JUEGO.md"
OUTPUT = ROOT / "Documentacion_Integracion_Admin_Juego.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
BLUE_DARK = "1F4D78"
MUTED = "60758A"
PALE = "F4F6F9"
BORDER = "D9E2EC"
PALE_BLUE = "E8EEF5"
WHITE = "FFFFFF"
GREEN = "227A55"


def set_cell_or_paragraph_shading(p, fill: str) -> None:
    p_pr = p._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_paragraph_borders(p, *, top=False, bottom=False, left=False, right=False,
                          color=BORDER, size="6", space="5") -> None:
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge, enabled in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        old = p_bdr.find(qn(f"w:{edge}"))
        if old is not None:
            p_bdr.remove(old)
        if enabled:
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), size)
            el.set(qn("w:space"), space)
            el.set(qn("w:color"), color)
            p_bdr.append(el)


def set_keep(p, *, next_=False, together=False) -> None:
    pf = p.paragraph_format
    pf.keep_with_next = next_
    pf.keep_together = together
    pf.widow_control = True


def set_lang(run, lang="es-GT") -> None:
    r_pr = run._r.get_or_add_rPr()
    lang_el = r_pr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        r_pr.append(lang_el)
    lang_el.set(qn("w:val"), lang)


def set_run_font(run, name: str, size: float | None = None, color: str | None = None,
                 bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    set_lang(run)


def add_field(run, field_code: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_el = OxmlElement("w:t")
    text_el.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text_el, end])


def configure_numbering(doc: Document) -> tuple[int, int]:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]

    def make_abstract(abstract_id: int, kind: str) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "multilevel")
        abstract.append(multi)
        for level in range(4):
            lvl = OxmlElement("w:lvl")
            lvl.set(qn("w:ilvl"), str(level))
            start = OxmlElement("w:start")
            start.set(qn("w:val"), "1")
            num_fmt = OxmlElement("w:numFmt")
            num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
            lvl_text = OxmlElement("w:lvlText")
            if kind == "bullet":
                lvl_text.set(qn("w:val"), "•" if level % 2 == 0 else "–")
            else:
                lvl_text.set(qn("w:val"), f"%{level + 1}.")
            lvl_jc = OxmlElement("w:lvlJc")
            lvl_jc.set(qn("w:val"), "left")
            p_pr = OxmlElement("w:pPr")
            tabs = OxmlElement("w:tabs")
            tab = OxmlElement("w:tab")
            tab.set(qn("w:val"), "num")
            left = 540 + level * 360
            tab.set(qn("w:pos"), str(left))
            tabs.append(tab)
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(left))
            ind.set(qn("w:hanging"), "270")
            p_pr.extend([tabs, ind])
            lvl.extend([start, num_fmt, lvl_text, lvl_jc, p_pr])
            if kind == "bullet":
                r_pr = OxmlElement("w:rPr")
                r_fonts = OxmlElement("w:rFonts")
                r_fonts.set(qn("w:ascii"), "Calibri")
                r_fonts.set(qn("w:hAnsi"), "Calibri")
                r_pr.append(r_fonts)
                lvl.append(r_pr)
            abstract.append(lvl)
        numbering.append(abstract)

    def make_num(num_id: int, abstract_id: int) -> None:
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abstract_id))
        num.append(ref)
        numbering.append(num)

    bullet_abs = max(abstract_ids, default=0) + 1
    decimal_abs = bullet_abs + 1
    bullet_num = max(num_ids, default=0) + 1
    decimal_num = bullet_num + 1
    make_abstract(bullet_abs, "bullet")
    make_abstract(decimal_abs, "decimal")
    make_num(bullet_num, bullet_abs)
    make_num(decimal_num, decimal_abs)
    return bullet_num, decimal_num


def apply_num(p, num_id: int, level: int = 0) -> None:
    p_pr = p._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def new_numbering_instance(doc: Document, base_num_id: int) -> int:
    numbering = doc.part.numbering_part.element
    base = next(
        item for item in numbering.findall(qn("w:num"))
        if int(item.get(qn("w:numId"))) == base_num_id
    )
    abstract_id = base.find(qn("w:abstractNumId")).get(qn("w:val"))
    new_id = max(int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), abstract_id)
    num.append(ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return new_id


INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")


def add_inline(p, text: str, *, base_size=11.0, base_color=None) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = p.add_run(text[cursor:match.start()])
            set_run_font(run, "Calibri", base_size, base_color)
        token = match.group(0)
        if token.startswith("`"):
            run = p.add_run(token[1:-1])
            set_run_font(run, "Consolas", max(base_size - 1, 8.5), BLUE_DARK)
            set_run_shading(run, PALE_BLUE)
        else:
            run = p.add_run(token[2:-2])
            set_run_font(run, "Calibri", base_size, base_color, bold=True)
        cursor = match.end()
    if cursor < len(text):
        run = p.add_run(text[cursor:])
        set_run_font(run, "Calibri", base_size, base_color)


def set_run_shading(run, fill: str) -> None:
    r_pr = run._r.get_or_add_rPr()
    shd = r_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        r_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def add_code_block(doc: Document, lines: list[str], language: str) -> None:
    if language == "mermaid":
        source = "\n".join(lines)
        if "sequenceDiagram" in source:
            raw = (
                "1. El académico crea la constelación/tema y sus estrellas/niveles.\n"
                "2. Selecciona una mecánica; la consola consulta su esquema y muestra los campos válidos.\n"
                "3. Completa contenido y recursos; el backend guarda y valida el borrador.\n"
                "4. El contenido pasa por revisión y aprobación.\n"
                "5. El publicador realiza la validación integral y genera un catálogo v2 inmutable.\n"
                "6. El juego consulta, valida, almacena y muestra la revisión publicada."
            )
        else:
            raw = (
                "Autoría: Consola académica web → API de autoría → Backend de contenido → "
                "PostgreSQL / almacenamiento multimedia.\n"
                "Publicación: Backend → Servicio de publicación → Catálogo JSON inmutable / CDN de recursos.\n"
                "Consumo: App Android → catálogo y recursos por HTTPS; App Android → API de progreso "
                "mediante eventos idempotentes."
            )
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.right_indent = Inches(0.08)
        p.paragraph_format.line_spacing = 1.15
        set_paragraph_borders(p, top=True, bottom=True, left=True, right=True, color=BLUE)
        set_cell_or_paragraph_shading(p, PALE_BLUE)
        r = p.add_run("FLUJO  ")
        set_run_font(r, "Calibri", 9, BLUE, bold=True)
        add_inline(p, raw, base_size=9.5, base_color=NAVY)
        set_keep(p, together=False)
        return

    for i, line in enumerate(lines or [""]):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.right_indent = Inches(0.08)
        p.paragraph_format.space_before = Pt(4 if i == 0 else 0)
        p.paragraph_format.space_after = Pt(4 if i == len(lines) - 1 else 0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        set_cell_or_paragraph_shading(p, PALE)
        set_paragraph_borders(
            p,
            top=i == 0,
            bottom=i == len(lines) - 1,
            left=True,
            right=True,
            color=BORDER,
            size="4",
            space="4",
        )
        r = p.add_run(line if line else " ")
        set_run_font(r, "Consolas", 8.2, BLUE_DARK)
        set_keep(p, next_=i < len(lines) - 1 and i < 3, together=False)


def set_style(style, *, font, size, color, bold=False, before=0, after=0, line=1.0):
    style.font.name = font
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.widow_control = True


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("CAZA ESTELAR  /  ESPECIFICACIÓN TÉCNICA")
    set_run_font(r, "Calibri", 9.5, BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Integración de la consola académica con Caza Estelar")
    set_run_font(r, "Calibri", 28, NAVY, bold=True)
    set_keep(p, next_=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("Contrato de contenido, publicación, sincronización y migración")
    set_run_font(r, "Calibri", 14, MUTED)

    for label, value in (
        ("ESTADO", "Propuesta técnica para implementación"),
        ("FECHA", "20 de agosto de 2026"),
        ("ALCANCE", "Consola web · Backend · Android/Kotlin/libGDX"),
        ("RESPONSABLE FUNCIONAL", "Equipo académico"),
        ("RESPONSABLE TÉCNICO", "Equipo de Caza Estelar"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"{label}  ")
        set_run_font(r, "Calibri", 8.5, BLUE, bold=True)
        r = p.add_run(value)
        set_run_font(r, "Calibri", 10.5, NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    set_cell_or_paragraph_shading(p, PALE_BLUE)
    set_paragraph_borders(p, top=True, bottom=True, left=True, right=True, color=BLUE)
    r = p.add_run("MODELO FUNCIONAL APROBADO\n")
    set_run_font(r, "Calibri", 9, BLUE, bold=True)
    r = p.add_run("Constelación / tema matemático  →  Estrella / nivel  →  Mecánica configurable")
    set_run_font(r, "Calibri", 12, NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(
        "Esta especificación define cómo el equipo académico podrá crear y publicar niveles "
        "matemáticos sin recompilar el juego. El desarrollo mantiene el control del código de "
        "las mecánicas; la consola administra su configuración, contenido, recursos y ciclo de aprobación."
    )
    set_run_font(r, "Calibri", 11, MUTED)
    p.paragraph_format.line_spacing = 1.25

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("Documento de referencia para producto, academia, backend y aplicación móvil")
    set_run_font(r, "Calibri", 9, GREEN, bold=True)
    p.add_run().add_break(WD_BREAK.PAGE)


def configure_document() -> tuple[Document, int, int]:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    sec.different_first_page_header_footer = True

    styles = doc.styles
    set_style(styles["Normal"], font="Calibri", size=11, color=NAVY, before=0, after=6, line=1.25)
    set_style(styles["Heading 1"], font="Calibri", size=16, color=BLUE, bold=True, before=18, after=10, line=1.05)
    set_style(styles["Heading 2"], font="Calibri", size=13, color=BLUE, bold=True, before=14, after=7, line=1.08)
    set_style(styles["Heading 3"], font="Calibri", size=12, color=BLUE_DARK, bold=True, before=10, after=5, line=1.1)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[style_name].paragraph_format.keep_with_next = True
        styles[style_name].paragraph_format.keep_together = True

    if "Code Caption" not in styles:
        code_caption = styles.add_style("Code Caption", WD_STYLE_TYPE.PARAGRAPH)
        set_style(code_caption, font="Calibri", size=8.5, color=BLUE, bold=True, before=6, after=2, line=1.0)

    header = sec.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    r = hp.add_run("CAZA ESTELAR  |  Integración administrativa")
    set_run_font(r, "Calibri", 8, MUTED, bold=True)
    r = hp.add_run("\tESPECIFICACIÓN v1.0")
    set_run_font(r, "Calibri", 8, MUTED)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    r = fp.add_run("Página ")
    set_run_font(r, "Calibri", 8, MUTED)
    field_run = fp.add_run()
    set_run_font(field_run, "Calibri", 8, MUTED)
    add_field(field_run, "PAGE")

    doc.core_properties.title = "Integración de la consola académica con Caza Estelar"
    doc.core_properties.subject = "Contrato técnico de contenido, publicación y sincronización"
    doc.core_properties.author = "Equipo de Caza Estelar"
    doc.core_properties.keywords = "Caza Estelar, Kotlin, libGDX, consola académica, niveles, matemáticas"
    doc.core_properties.comments = "Especificación técnica v1.0"

    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    bullet_num, decimal_num = configure_numbering(doc)
    return doc, bullet_num, decimal_num


def add_markdown_content(doc: Document, bullet_num: int, decimal_num: int, md: str) -> None:
    lines = md.splitlines()
    # The cover replaces the Markdown title and four metadata lines.
    start = next(i for i, line in enumerate(lines) if line.startswith("## 1. "))
    i = start
    code_language: str | None = None
    code_lines: list[str] = []
    numbered_active = False
    current_decimal_num = decimal_num
    while i < len(lines):
        line = lines[i]
        if code_language is not None:
            if line.startswith("```"):
                add_code_block(doc, code_lines, code_language)
                code_language = None
                code_lines = []
                numbered_active = False
            else:
                code_lines.append(line)
            i += 1
            continue

        fence = re.match(r"^```(.*)$", line)
        if fence:
            code_language = fence.group(1).strip().lower()
            code_lines = []
            numbered_active = False
            i += 1
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1)) - 1
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            add_inline(p, heading.group(2), base_size={1: 16, 2: 13, 3: 12}[min(level, 3)], base_color={1: BLUE, 2: BLUE, 3: BLUE_DARK}[min(level, 3)])
            for run in p.runs:
                run.bold = True
            numbered_active = False
            i += 1
            continue

        bullet = re.match(r"^(\s*)[-*]\s+(.+)$", line)
        if bullet:
            level = min(len(bullet.group(1)) // 2, 3)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            apply_num(p, bullet_num, level)
            add_inline(p, bullet.group(2))
            set_keep(p)
            numbered_active = False
            i += 1
            continue

        numbered = re.match(r"^(\s*)\d+\.\s+(.+)$", line)
        if numbered:
            if not numbered_active:
                current_decimal_num = new_numbering_instance(doc, decimal_num)
                numbered_active = True
            level = min(len(numbered.group(1)) // 2, 3)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.25
            apply_num(p, current_decimal_num, level)
            add_inline(p, numbered.group(2))
            set_keep(p)
            i += 1
            continue

        if line.strip():
            p = doc.add_paragraph()
            add_inline(p, line.strip())
            set_keep(p)
            numbered_active = False
        i += 1

    if code_language is not None:
        add_code_block(doc, code_lines, code_language)


def preset_audit(doc: Document) -> None:
    sec = doc.sections[0]
    expected = {
        "page_width": Inches(8.5),
        "page_height": Inches(11),
        "top_margin": Inches(1),
        "bottom_margin": Inches(1),
        "left_margin": Inches(1),
        "right_margin": Inches(1),
    }
    for attr, value in expected.items():
        actual = getattr(sec, attr)
        if abs(actual - value) > 10:
            raise AssertionError(f"Preset mismatch: {attr}={actual}, expected {value}")
    assert doc.styles["Normal"].font.name == "Calibri"
    assert doc.styles["Normal"].font.size == Pt(11)
    assert doc.styles["Heading 1"].font.size == Pt(16)
    assert doc.styles["Heading 2"].font.size == Pt(13)
    assert doc.styles["Heading 3"].font.size == Pt(12)
    numbering = doc.part.numbering_part.element
    assert len(numbering.findall(qn("w:num"))) >= 2


def main() -> None:
    source = SOURCE.read_text(encoding="utf-8")
    doc, bullet_num, decimal_num = configure_document()
    add_cover(doc)
    add_markdown_content(doc, bullet_num, decimal_num, source)
    preset_audit(doc)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")
    print(f"Paragraphs: {len(doc.paragraphs)}")


if __name__ == "__main__":
    main()
