from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK


ROOT = Path(r"C:\Users\explorax\Documents\Caza Estelar\Nueva carpeta")
OUT = ROOT / "Propuesta_Caza_Estelar.docx"
IMG = ROOT / "assets" / "images"
BRAND = ROOT / "assets" / "branding"

NAVY = "07152E"
BLUE = "185ADB"
CYAN = "18BFE8"
PURPLE = "6D4BE8"
GOLD = "F5B942"
INK = "18233A"
MUTED = "5B6780"
LIGHT = "F2F6FC"
PALE_BLUE = "EAF6FF"
PALE_PURPLE = "F3EEFF"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths_dxa)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths_dxa[idx]))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_font(run, size=None, bold=None, color=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_font(run, 9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_h1(doc, text):
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.keep_with_next = True
    return p


def add_h2(doc, text):
    p = doc.add_paragraph(text, style="Heading 2")
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, 11, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2, 11, color=INK)
    else:
        r = p.add_run(text)
        set_font(r, 11, color=INK)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    r = p.add_run(text)
    set_font(r, 10.5, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    r = p.add_run(text)
    set_font(r, 10.5, color=INK)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, accent=CYAN):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper() + "  ")
    set_font(r, 9.5, bold=True, color=accent)
    r = p.add_run(text)
    set_font(r, 10.5, bold=True, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc, path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", path.stem)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    cp.paragraph_format.keep_with_next = True
    r = cp.add_run(caption)
    set_font(r, 9, italic=True, color=MUTED)


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    set_table_geometry(table, [2700, 6660])
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        set_cell_shading(cells[0], NAVY if i % 2 == 0 else BLUE)
        set_cell_shading(cells[1], LIGHT)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[0].paragraphs[0]
        r = p.add_run(label)
        set_font(r, 10, bold=True, color=WHITE)
        p = cells[1].paragraphs[0]
        r = p.add_run(value)
        set_font(r, 10, color=INK)
    set_table_geometry(table, [2700, 6660])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_font(r, 9.5, bold=True, color=WHITE)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_shading(cells[i], WHITE if ridx % 2 == 0 else LIGHT)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_font(r, 9.2, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.75)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.25
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for style_name, size, color, before, after in (
    ("Heading 1", 17, NAVY, 16, 8),
    ("Heading 2", 13, BLUE, 11, 5),
    ("Heading 3", 11.5, PURPLE, 8, 4),
):
    st = styles[style_name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

for list_name in ("List Bullet", "List Bullet 2", "List Number"):
    st = styles[list_name]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.208

header = sec.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = hp.add_run("EXPLORAX  |  CAZA ESTELAR")
set_font(r, 8.5, bold=True, color=MUTED)
footer = sec.footer
add_page_number(footer.paragraphs[0])

# Cover - proposal_centerpiece
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PROPUESTA DE PROYECTO EDUCATIVO")
set_font(r, 10, bold=True, color=CYAN)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Caza Estelar")
set_font(r, 30, bold=True, color=NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
r = p.add_run("Aprendizaje interactivo de matemáticas para niños de 6 a 8 años")
set_font(r, 14, color=BLUE)
add_figure(doc, BRAND / "explorax-caza-estelar-intro-hd.png", 6.35, "Identidad visual conceptual de Caza Estelar")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
r = p.add_run("Propuesta funcional, académica y técnica")
set_font(r, 11, bold=True, color=NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Versión de presentación · Agosto de 2026")
set_font(r, 9.5, color=MUTED)

doc.add_page_break()

# 1
add_h1(doc, "1. Resumen ejecutivo")
add_callout(doc, "Idea central", "Convertir el refuerzo de matemáticas en una aventura visual donde cada logro devuelve el brillo a una galaxia apagada.")
add_body(doc, "Caza Estelar es una experiencia de aprendizaje interactivo, principalmente educativa, diseñada para niños de 6 a 8 años. Su uso se orienta a colegios y escuelas de Guatemala y Estados Unidos, con contenido matemático alineado al currículo vigente de Guatemala y disponible de forma equivalente en español e inglés.")
add_body(doc, "El proyecto propone una experiencia narrativa en la que Tojil guía al alumno en la búsqueda de estrellas perdidas. Cada constelación representa un tema matemático; sus estrellas son actividades que el alumno completa para recuperar progresivamente el cielo y transformar el observatorio.")
add_key_value_table(doc, [
    ("Público", "Niños de 6 a 8 años en contextos escolares e institucionales."),
    ("Enfoque", "Educativo con recursos de entretenimiento, exploración y recompensa visual."),
    ("Materia", "Matemáticas, organizadas por grado y conforme al currículo de Guatemala."),
    ("Plataforma", "Android para teléfonos y tabletas, exclusivamente horizontal."),
    ("Acceso", "Gratuito en Google Play y distribución adicional mediante APK institucional."),
])
add_h2(doc, "Objetivo de la propuesta")
add_body(doc, "Presentar una base común para dirección, desarrollo, diseño, QA y equipo académico. El documento define la visión, el funcionamiento esperado y los principales compromisos técnicos, conservando como pendientes las decisiones que dependen de contenido final, assets o evaluación arquitectónica.")

doc.add_page_break()

# 2
add_h1(doc, "2. Oportunidad educativa")
add_figure(doc, IMG / "explorax-observatory-login-v2.png", 6.15, "Ilustración conceptual del observatorio y el aprendizaje visual colaborativo")
add_h2(doc, "Planteamiento")
add_body(doc, "Caza Estelar busca complementar la clase, no sustituirla. La aplicación transforma ejercicios matemáticos aprobados por un equipo académico en experiencias breves, interactivas y visuales que ayudan a practicar los contenidos del grado del alumno.")
add_bullet(doc, "Asignación automática del contenido según el grado registrado en la cuenta.")
add_bullet(doc, "Actividades de selección, conteo, ordenamiento, operaciones, arrastre y otras mecánicas 2D.")
add_bullet(doc, "Bancos de ejercicios con variantes y orden cambiante para reducir la memorización.")
add_bullet(doc, "Acompañamiento positivo de Tojil en lugar de mensajes directos de error.")
add_bullet(doc, "Repetición ilimitada para practicar y mejorar el resultado.")
add_callout(doc, "Principio pedagógico", "El porcentaje representa el nivel de autonomía con el que el alumno resuelve la actividad, no solo el hecho de llegar finalmente a la respuesta correcta.", fill=PALE_PURPLE, accent=PURPLE)

doc.add_page_break()

# 3
add_h1(doc, "3. Universo narrativo y experiencia")
add_figure(doc, IMG / "login-characters-terrace-1920.png", 6.15, "Elenco conceptual de Explorax en la búsqueda de las constelaciones perdidas")
add_body(doc, "La historia comienza con la explosión natural y visual de dos estrellas. El acontecimiento apaga las constelaciones de la galaxia y da origen a la misión. Tojil, personaje principal y guía, se dirige al alumno por su nombre y lo acompaña en la recuperación del cielo.")
add_h2(doc, "Progresión narrativa")
add_number(doc, "El alumno entra al observatorio y descubre que las constelaciones han perdido su brillo.")
add_number(doc, "Selecciona una constelación disponible y elige libremente una de sus actividades.")
add_number(doc, "Cada actividad aprobada recupera una estrella del tema matemático.")
add_number(doc, "Al completar todas las estrellas, la constelación vuelve a brillar y el observatorio cambia de forma permanente.")
add_number(doc, "Al recuperar todas las constelaciones, el cielo se ilumina por completo y el explorador recibe una felicitación final.")
add_body(doc, "Los demás personajes de Explorax —Chnám, Ik’ar, Ulew, Rum, Xam, Ya’ y Ja’— actuarán como aliados secundarios. Podrán apoyar ciertos temas, aparecer en el observatorio al terminarlos y, en algunos casos, interactuar mediante globos de texto.")

doc.add_page_break()

# 4
add_h1(doc, "4. Modelo de constelaciones y actividades")
add_figure(doc, IMG / "explorax-constellation-menu-v3.png", 6.15, "Representación conceptual del mapa de constelaciones y sus actividades")
add_key_value_table(doc, [
    ("Constelación", "Representa un tema matemático completo."),
    ("Estrella", "Representa una actividad del tema."),
    ("Cantidad", "Aproximadamente seis actividades por constelación."),
    ("Intento", "Presenta cuatro o cinco ejercicios tomados de un banco mayor."),
    ("Aprobación", "Cada actividad requiere dos estrellas, equivalentes a 60 % o más."),
    ("Recuperación", "La constelación se recupera al aprobar todas sus actividades."),
])
add_h2(doc, "Navegación y desbloqueos")
add_body(doc, "El recorrido será no lineal. Dentro de una constelación disponible, el alumno podrá escoger cualquier actividad. Los temas independientes estarán abiertos sin un orden obligatorio; una constelación complementaria se desbloqueará únicamente cuando se haya recuperado por completo su requisito previo.")
add_h2(doc, "Variación de contenido")
add_body(doc, "El equipo académico entregará planteamientos, variantes y respuestas correctas. El sistema evitará repetir ejercicios para un alumno hasta agotar las variantes disponibles y después iniciará un nuevo ciclo de selección con dificultad similar.")

doc.add_page_break()

# 5
add_h1(doc, "5. Evaluación, ayudas y recompensa")
add_matrix(doc,
    ["Desempeño final", "Recompensa", "Resultado"],
    [
        ("0–29 %", "0 estrellas", "Actividad no aprobada"),
        ("30–59 %", "1 estrella", "Actividad no aprobada"),
        ("60–89 %", "2 estrellas", "Actividad aprobada"),
        ("90–100 %", "3 estrellas", "Dominio destacado"),
    ],
    [3000, 2500, 3860]
)
add_h2(doc, "Puntuación por autonomía")
add_matrix(doc,
    ["Resolución", "Valor del ejercicio"],
    [
        ("Correcta sin ayuda", "100 %"),
        ("Correcta después de una ayuda", "75 %"),
        ("Varias ayudas o solución mostrada", "50 %"),
    ],
    [6500, 2860]
)
add_h2(doc, "Secuencia de apoyo de Tojil")
add_number(doc, "Después de tres respuestas no acertadas consecutivas, Tojil muestra la primera ayuda.")
add_number(doc, "Una nueva respuesta no acertada activa una segunda ayuda.")
add_number(doc, "Si el alumno vuelve a intentarlo sin acertar, Tojil muestra la solución correcta.")
add_body(doc, "Las respuestas se presentan con lenguaje de apoyo —por ejemplo, «No estoy seguro»— y no como una sanción. Cada ejercicio tiene el mismo valor; cuando contiene varias partes, el valor se divide automáticamente de forma equitativa. El porcentaje entero redondeado aparece únicamente al terminar.")
add_callout(doc, "Reintentos", "No existe límite. El mejor porcentaje queda como resultado principal, pero la bitácora conserva todos los intentos, su tiempo y número.")

doc.add_page_break()

# 6
add_h1(doc, "6. Ecosistema institucional y bitácora")
add_body(doc, "Caza Estelar se integrará con una web independiente que centraliza usuarios válidos y resultados de varios juegos del entorno Explorax. El administrador general crea las cuentas desde listados institucionales y el sistema genera usuarios y contraseñas.")
add_h2(doc, "Roles y alcance")
add_matrix(doc,
    ["Rol", "Alcance principal"],
    [
        ("Docente", "Nombres completos e información académica de sus secciones asignadas."),
        ("Director", "Todas las secciones de su colegio y exportación de reportes."),
        ("Administrador general", "Todos los colegios y países; cuentas, auditoría y fecha de nacimiento."),
        ("Alumno", "Acceso individual al juego mediante usuario y contraseña."),
    ],
    [2600, 6760]
)
add_h2(doc, "Resultados y reportes")
add_bullet(doc, "Envío de tiempo, porcentaje y número de intento; las respuestas individuales no se envían.")
add_bullet(doc, "Filtros por fecha, grado, sección y actividad; exportación académica en PDF y Excel.")
add_bullet(doc, "Promedio del grupo calculado con el mejor porcentaje de cada alumno; quien no participa cuenta como 0 %.")
add_bullet(doc, "Identificación de actividades incompletas mediante estado y porcentaje parcial.")
add_bullet(doc, "Reporte de credenciales separado y exclusivo del administrador general.")
add_callout(doc, "Persistencia", "El punto exacto para reanudar una actividad incompleta permanece localmente en el dispositivo; la web solo conoce el estado y el porcentaje alcanzado.", fill=PALE_PURPLE, accent=PURPLE)

doc.add_page_break()

# 7
add_h1(doc, "7. Cuentas, privacidad y continuidad")
add_h2(doc, "Datos del alumno")
add_body(doc, "Cada cuenta incluye nombre real, fecha de nacimiento, grado, sección, colegio y país. La edad se calcula automáticamente desde la fecha de nacimiento. No se utilizarán alias y solo el administrador general podrá consultar la fecha de nacimiento.")
add_h2(doc, "Reglas de sesión")
add_bullet(doc, "La sesión permanece abierta al cerrar y volver a abrir la aplicación.")
add_bullet(doc, "Al cerrar sesión se eliminan contraseña, sesión e intentos incompletos; puede conservarse el usuario como sugerencia.")
add_bullet(doc, "La pantalla permite olvidar el usuario sugerido en dispositivos compartidos.")
add_bullet(doc, "El administrador general puede bloquear cuentas y cerrar sesiones remotamente.")
add_bullet(doc, "Una cuenta reactivada conserva credenciales, historial y progreso terminado.")
add_h2(doc, "Privacidad institucional")
add_body(doc, "La institución será responsable de obtener la autorización de padres o encargados. Los datos se conservan como historial hasta que el administrador general ejecute una eliminación definitiva. Al eliminar una cuenta, sus resultados desaparecen de promedios y reportes; los eventos de auditoría se conservan sin datos identificables.")
add_callout(doc, "Auditoría", "La web registra inicios y cierres de sesión, bloqueos y desbloqueos. Solo el administrador general consulta el registro.")

doc.add_page_break()

# 8
add_h1(doc, "8. Arquitectura técnica propuesta")
add_callout(doc, "Decisión objetivo", "Reescribir el prototipo actual de Java/libGDX completamente en Kotlin.")
add_h2(doc, "Aplicación Android")
add_bullet(doc, "Kotlin como lenguaje objetivo y Jetpack Compose para la experiencia Android.")
add_bullet(doc, "Evaluación pendiente entre Compose para interfaz + libGDX para actividades, o Compose/Canvas para toda la aplicación.")
add_bullet(doc, "Experiencia 2D con numerosos elementos animados en observatorio, constelaciones y actividades.")
add_bullet(doc, "Android 8.0/API 26 como versión mínima; teléfonos y tabletas en orientación horizontal.")
add_h2(doc, "Servicios y datos")
add_bullet(doc, "Backend exclusivo para Caza Estelar, como proyecto separado de los demás juegos.")
add_bullet(doc, "Base de datos compartida con la web/bitácora.")
add_bullet(doc, "Infraestructura orientada inicialmente a AWS.")
add_bullet(doc, "API compartida para autenticación e intercambio de resultados; formato técnico aún pendiente.")
add_h2(doc, "Integridad de resultados")
add_body(doc, "Cada intento tendrá un identificador único para evitar duplicados. Si la API no recibe un resultado, la aplicación lo conserva localmente y reintenta el envío cuando vuelve la conexión. Durante una desconexión, la actividad y el temporizador se pausan.")

doc.add_page_break()

# 9
add_h1(doc, "9. Calidad, rendimiento y distribución")
add_key_value_table(doc, [
    ("Rendimiento", "Objetivo de 60 FPS en la matriz mínima de dispositivos."),
    ("Inicio", "Menos de 10 segundos hasta mostrar el acceso."),
    ("Actividad", "Máximo 2 segundos para abrir en condiciones normales."),
    ("Tamaño", "Google Play y APK no deben superar 500 MB."),
    ("Calidad visual", "Reducción automática de efectos cuando no se sostengan 60 FPS."),
])
add_h2(doc, "Criterios de lanzamiento")
add_bullet(doc, "No se publica una versión con errores visuales conocidos, aunque sean menores.")
add_bullet(doc, "Cualquier defecto que afecte cuentas, progreso, porcentajes o resultados bloquea el lanzamiento.")
add_bullet(doc, "QA valida todas las actividades en español e inglés.")
add_bullet(doc, "El equipo académico revisa la actividad implementada además del contenido original.")
add_bullet(doc, "Pruebas en teléfonos, tabletas físicas y emuladores, incluyendo segundo plano, desconexión y cierres.")
add_bullet(doc, "Pruebas automatizadas para porcentajes, estrellas, desbloqueos y envío de resultados.")
add_h2(doc, "Entrega")
add_body(doc, "El juego será gratuito en Google Play y también se distribuirá mediante APK institucional. Existirán ambientes separados de desarrollo, pruebas y producción. La dirección del proyecto autorizará cada publicación.")

doc.add_page_break()

# 10
add_h1(doc, "10. Alcance de la primera propuesta")
add_h2(doc, "Incluido")
add_bullet(doc, "Aplicación Android horizontal para teléfonos y tabletas.")
add_bullet(doc, "Cuentas individuales, autenticación y sesión persistente.")
add_bullet(doc, "Constelaciones matemáticas, actividades variables, ayudas y recompensas.")
add_bullet(doc, "Narrativa de Tojil, personajes aliados, observatorio evolutivo y final general.")
add_bullet(doc, "Integración con la bitácora, reportes institucionales y control por roles.")
add_bullet(doc, "Soporte bilingüe equivalente, música y efectos sin dependencia del sonido.")
add_h2(doc, "Pendiente o diferido")
add_bullet(doc, "Grados, temas matemáticos y cantidad final de constelaciones.")
add_bullet(doc, "Assets definitivos, animaciones individuales y guías detalladas de personajes.")
add_bullet(doc, "Elección formal Compose/libGDX frente a Compose/Canvas mediante comparación y prototipo técnico.")
add_bullet(doc, "Tecnología concreta del backend, servicios AWS y contrato definitivo de la API.")
add_bullet(doc, "Cambio obligatorio de contraseña y formato del reporte de credenciales.")
add_bullet(doc, "Diseño de paquetes, caché y descarga posterior de assets.")
add_bullet(doc, "Políticas operativas detalladas de respaldo, retención y respuesta a incidentes.")
add_callout(doc, "Estado", "La orientación funcional, educativa, narrativa e institucional está definida. Los puntos abiertos dependen de decisiones técnicas, contenido académico o producción visual futura.", fill=PALE_BLUE, accent=BLUE)

doc.add_page_break()

# 11
add_h1(doc, "11. Ruta de trabajo propuesta")
add_number(doc, "Validar la propuesta con dirección, equipo académico, diseño, desarrollo y QA.")
add_number(doc, "Definir grados, mapa curricular y primera constelación con aproximadamente seis actividades.")
add_number(doc, "Recibir assets finales y guías de personajes desde Explorax.")
add_number(doc, "Ejecutar un prototipo comparativo para seleccionar Compose/libGDX o Compose/Canvas.")
add_number(doc, "Definir backend, servicios AWS, esquema compartido y contrato de API.")
add_number(doc, "Construir una vertical completa: acceso, una constelación, bitácora y reporte.")
add_number(doc, "Validar rendimiento, pedagogía, bilingüismo y experiencia en la matriz de dispositivos.")
add_number(doc, "Promover por desarrollo, pruebas y producción con aprobación de dirección.")
add_h2(doc, "Criterio de éxito")
add_body(doc, "La propuesta será exitosa si el alumno puede reforzar matemáticas con autonomía creciente, si el docente obtiene resultados comprensibles y si la experiencia mantiene un rendimiento fluido y una narrativa motivadora en dispositivos institucionales reales.")
add_callout(doc, "Visión final", "Cada ejercicio recupera una estrella; cada tema devuelve una constelación; cada constelación transforma el observatorio hasta volver a encender todo el cielo.", fill=PALE_PURPLE, accent=PURPLE)
add_body(doc, "Nota sobre imágenes: las ilustraciones incluidas en esta propuesta son recursos conceptuales provisionales. Los assets finales serán entregados y aprobados posteriormente por Explorax.")

doc.core_properties.title = "Propuesta Caza Estelar"
doc.core_properties.subject = "Propuesta funcional, académica y técnica"
doc.core_properties.author = "Explorax"
doc.core_properties.keywords = "Caza Estelar, Explorax, matemáticas, educación, Android"

doc.save(OUT)
print(OUT)
