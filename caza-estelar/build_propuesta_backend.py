from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\explorax\Documents\Caza Estelar\Nueva carpeta")
OUT = ROOT / "Propuesta_Backend_Caza_Estelar.docx"
DIAGRAM = ROOT / "arquitectura_backend_caza_estelar.png"
LOGO = ROOT / "assets" / "branding" / "explorax-caza-estelar-intro-hd.png"

NAVY = "07152E"
BLUE = "185ADB"
CYAN = "18BFE8"
PURPLE = "6D4BE8"
GOLD = "F5B942"
INK = "18233A"
MUTED = "5B6780"
LIGHT = "F2F6FC"
PALE_BLUE = "EAF6FF"
PALE_PURPLE = "F3EEFF"
PALE_GOLD = "FFF7E3"
WHITE = "FFFFFF"
GREEN = "17845B"
RED = "B23A48"


def font(size=34, bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf"),
    ]
    for p in candidates:
        if p.exists():
            return ImageFont.truetype(str(p), size)
    return ImageFont.load_default()


def make_architecture_diagram():
    img = Image.new("RGB", (1800, 1020), "#F7FAFF")
    d = ImageDraw.Draw(img)
    title = font(48, True)
    body = font(29, False)
    small = font(24, False)
    d.text((70, 42), "Arquitectura de referencia - Caza Estelar", font=title, fill="#07152E")

    boxes = [
        ((70, 180, 390, 430), "Aplicación Android\nKotlin", "Sesión, actividades,\ncola local y reintentos", "#EAF6FF", "#185ADB"),
        ((520, 180, 880, 430), "Capa de API", "HTTPS + JSON propuesto\nJWT, validación y límites", "#F3EEFF", "#6D4BE8"),
        ((1010, 150, 1510, 460), "Backend Caza Estelar", "Monolito modular\nAutenticación integrada\nProgreso y resultados\nVersiones y auditoría", "#EAF8F1", "#17845B"),
        ((1070, 620, 1530, 890), "Base compartida", "Esquema propio del juego\nVistas para bitácora\nBackups y cifrado", "#FFF7E3", "#C27A00"),
        ((510, 620, 900, 890), "Servicios AWS", "Secrets Manager\nCloudWatch\nS3 / artefactos\nIaC y ambientes", "#F5F7FA", "#5B6780"),
    ]
    for (x1, y1, x2, y2), heading, text, fill, outline in boxes:
        d.rounded_rectangle((x1, y1, x2, y2), radius=24, fill=fill, outline=outline, width=5)
        lines = heading.split("\n")
        y = y1 + 32
        for line in lines:
            d.text((x1 + 28, y), line, font=font(32, True), fill="#07152E")
            y += 42
        y += 10
        for line in text.split("\n"):
            d.text((x1 + 28, y), line, font=small, fill="#263653")
            y += 34

    def arrow(a, b, color="#185ADB"):
        d.line((a, b), fill=color, width=10)
        bx, by = b
        ax, ay = a
        if abs(bx - ax) > abs(by - ay):
            pts = [(bx, by), (bx - 24 if bx > ax else bx + 24, by - 16), (bx - 24 if bx > ax else bx + 24, by + 16)]
        else:
            pts = [(bx, by), (bx - 16, by - 24 if by > ay else by + 24), (bx + 16, by - 24 if by > ay else by + 24)]
        d.polygon(pts, fill=color)

    arrow((390, 305), (520, 305))
    arrow((880, 305), (1010, 305))
    arrow((1270, 460), (1270, 620), "#17845B")
    arrow((1010, 390), (850, 620), "#5B6780")
    d.text((70, 940), "Principio: la aplicación móvil nunca accede directamente a la base de datos.", font=body, fill="#B23A48")
    img.save(DIAGRAM, quality=95)


def set_font(run, size=None, bold=None, color=None, italic=None, name="Arial"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), str(indent))
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths[i]))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def mark_header(row):
    trPr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trPr.append(node)


def add_page_number(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Página ")
    set_font(r, 9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


def h1(doc, text):
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.keep_with_next = True
    return p


def h2(doc, text):
    p = doc.add_paragraph(text, style="Heading 2")
    p.paragraph_format.keep_with_next = True
    return p


def body(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, 10.7, color=INK)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, 10.3, color=INK)
    return p


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_font(r, 10.3, color=INK)
    return p


def callout(doc, status, title, text):
    palette = {
        "CONFIRMADO": (PALE_BLUE, BLUE),
        "RECOMENDADO": (PALE_PURPLE, PURPLE),
        "PENDIENTE": (PALE_GOLD, "A66B00"),
        "RIESGO": ("FCECEF", RED),
    }
    fill, accent = palette[status]
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(status + "  ")
    set_font(r, 9.2, bold=True, color=accent)
    r = p.add_run(title + ". ")
    set_font(r, 10.5, bold=True, color=INK)
    r = p.add_run(text)
    set_font(r, 10.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, 9.2, bold=True, color=WHITE)
    mark_header(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_shading(cells[i], WHITE if ridx % 2 == 0 else LIGHT)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            set_font(r, 8.9, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_picture(doc, path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", path.stem)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    r = cp.add_run(caption)
    set_font(r, 8.8, italic=True, color=MUTED)


make_architecture_diagram()
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.78)
sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.7)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for name, size, color, before, after in (
    ("Heading 1", 16, NAVY, 14, 7),
    ("Heading 2", 12.5, BLUE, 10, 5),
    ("Heading 3", 11.5, PURPLE, 8, 4),
):
    st = doc.styles[name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

for name in ("List Bullet", "List Number"):
    st = doc.styles[name]
    st.font.name = "Arial"
    st.font.size = Pt(10.3)
    st.paragraph_format.line_spacing = 1.167

hp = sec.header.paragraphs[0]
r = hp.add_run("CAZA ESTELAR  |  PROPUESTA DE BACKEND")
set_font(r, 8.4, bold=True, color=MUTED)
add_page_number(sec.footer.paragraphs[0])

# Cover - memo masthead
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
r = p.add_run("PROPUESTA TÉCNICA")
set_font(r, 10, bold=True, color=CYAN)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(3)
r = p.add_run("Backend de Caza Estelar")
set_font(r, 28, bold=True, color=NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
r = p.add_run("Arquitectura, datos, seguridad, integración y operación en AWS")
set_font(r, 13.5, color=BLUE)
add_picture(doc, LOGO, 6.2, "Identidad conceptual de Caza Estelar")
matrix(doc, ["Documento", "Alcance", "Estado"], [
    ("Propuesta de backend", "Servicio exclusivo para Caza Estelar e integración con bitácora", "Para evaluación"),
    ("Infraestructura", "AWS con desarrollo, pruebas y producción", "Orientación confirmada"),
    ("Tecnología", "Kotlin/JVM, contenedores o serverless", "Pendiente de decisión"),
], [2600, 4800, 1960])
body(doc, "Fecha de preparación: agosto de 2026. Audiencia: dirección del proyecto, desarrollo, arquitectura, QA y responsables de la web/bitácora.")

doc.add_page_break()

h1(doc, "1. Decisión ejecutiva propuesta")
callout(doc, "RECOMENDADO", "Arquitectura inicial", "Implementar un monolito modular exclusivo para Caza Estelar, desplegado en AWS, con acceso controlado al esquema del juego dentro de la base compartida y una API versionada como única entrada desde la aplicación móvil.")
body(doc, "La arquitectura debe comenzar simple, observable y segura. El volumen inicial todavía no está definido y varios componentes —tecnología del backend, formato final de API y servicios AWS concretos— permanecen abiertos. Por ello se propone una base que permita validar el producto sin introducir microservicios prematuros.")
h2(doc, "Resultados esperados")
bullet(doc, "Autenticar cuentas generadas desde los listados institucionales y mantener sesiones persistentes.")
bullet(doc, "Recibir resultados, intentos incompletos y eventos de versión sin duplicar información.")
bullet(doc, "Proteger los datos de menores mediante permisos por rol y separación de responsabilidades.")
bullet(doc, "Integrarse con la web/bitácora usando una base compartida sin permitir escrituras descontroladas entre sistemas.")
bullet(doc, "Operar en desarrollo, pruebas y producción con diagnósticos, backups y despliegues repetibles.")
callout(doc, "CONFIRMADO", "Límite del servicio", "El backend pertenece a Caza Estelar y no prestará servicios directamente a los demás juegos de Explorax.")

doc.add_page_break()

h1(doc, "2. Requisitos confirmados")
matrix(doc, ["Área", "Requisito confirmado"], [
    ("Usuarios", "Cuentas individuales creadas por el administrador general desde listados institucionales."),
    ("Sesión", "Usuario y contraseña; sesión persistente; bloqueo y cierre remoto."),
    ("Resultados", "Tiempo, porcentaje y número de intento; todos los intentos se conservan."),
    ("Idempotencia", "Cada intento tiene identificador único y los reenvíos no crean duplicados."),
    ("Incompletos", "La web recibe estado y porcentaje parcial; el punto exacto permanece local."),
    ("Reportes", "La versión reciente alimenta reportes; datos anteriores se conservan históricamente."),
    ("Datos", "La base es compartida con la web/bitácora."),
    ("Nube", "AWS como orientación inicial."),
], [2200, 7160])
h2(doc, "Restricciones de producto que afectan al backend")
bullet(doc, "El juego es totalmente en línea; una desconexión pausa la actividad y el temporizador.")
bullet(doc, "Los resultados no confirmados por la API se conservan localmente y se reenvían.")
bullet(doc, "Una versión anterior no puede continuar jugando cuando existe una actualización obligatoria.")
bullet(doc, "Los resultados de una cuenta eliminada deben desaparecer de promedios y reportes, manteniendo auditoría anonimizada.")
bullet(doc, "El historial puede trasladarse cuando el alumno cambia de colegio.")

doc.add_page_break()

h1(doc, "3. Arquitectura de referencia")
add_picture(doc, DIAGRAM, 6.45, "Flujo propuesto entre aplicación, API, backend, servicios AWS y base compartida")
h2(doc, "Principios")
bullet(doc, "La aplicación nunca se conecta directamente a la base de datos.")
bullet(doc, "La API valida identidad, versión mínima, formato e idempotencia antes de ejecutar cambios.")
bullet(doc, "El backend es dueño de las reglas de persistencia de Caza Estelar.")
bullet(doc, "La bitácora consume vistas o contratos estables; no depende de tablas internas sin control.")
bullet(doc, "La infraestructura se define como código y se reproduce por ambiente.")
callout(doc, "PENDIENTE", "Formato de API", "HTTPS con JSON y tokens JWT es la opción recomendada, pero todavía requiere aprobación conjunta con el equipo responsable de la API.")

doc.add_page_break()

h1(doc, "4. Diseño interno del backend")
body(doc, "Se recomienda un monolito modular: un único despliegue con módulos bien delimitados, contratos internos claros y una sola responsabilidad operativa. Esta estructura reduce complejidad al inicio y permite separar servicios únicamente cuando el volumen o la organización lo justifiquen.")
matrix(doc, ["Módulo", "Responsabilidad"], [
    ("Identidad e integración", "Validar identidad, sesión, bloqueo y pertenencia institucional."),
    ("Perfiles académicos", "Consultar grado, sección, colegio y país sin exponer datos innecesarios."),
    ("Catálogo y versiones", "Registrar versión del juego, actividades y compatibilidad de resultados."),
    ("Sesiones de actividad", "Crear intentos, marcar incompletos y aceptar finalizaciones."),
    ("Resultados", "Persistir tiempo, porcentaje, estrellas, mejor intento e historial."),
    ("Progresión", "Resolver desbloqueos, constelaciones recuperadas y progreso general."),
    ("Auditoría", "Registrar accesos, bloqueos, desbloqueos, traslados y eliminaciones."),
    ("Integración de bitácora", "Exponer vistas o eventos listos para reportes institucionales."),
], [2700, 6660])
h2(doc, "Elección tecnológica")
callout(doc, "RECOMENDADO", "Prueba de concepto", "Comparar Kotlin/JVM con Ktor o Spring Boot en contenedores frente a funciones Lambda basadas en JVM. Medir tiempo de desarrollo, latencia, consumo, conexión a la base compartida y facilidad de operación.")
body(doc, "Para un servicio con base relacional compartida, reglas de dominio y múltiples módulos, un contenedor Kotlin/JVM en ECS Fargate ofrece una ruta directa y predecible. Lambda sigue siendo viable para cargas intermitentes, pero exige evaluar arranque en frío y administración de conexiones; RDS Proxy puede reducir el costo de abrir y mantener conexiones.")

doc.add_page_break()

h1(doc, "5. Contrato de API propuesto")
matrix(doc, ["Operación", "Ruta conceptual", "Notas"], [
    ("Iniciar sesión", "POST /v1/auth/login", "Usuario, contraseña, dispositivo y versión."),
    ("Renovar sesión", "POST /v1/auth/refresh", "Renovación silenciosa de sesión persistente."),
    ("Cerrar sesión", "POST /v1/auth/logout", "Revoca sesión del dispositivo."),
    ("Comprobar versión", "GET /v1/app/version", "Devuelve versión mínima y canales de actualización."),
    ("Crear intento", "POST /v1/attempts", "Genera o acepta attemptId único."),
    ("Actualizar incompleto", "PATCH /v1/attempts/{id}", "Estado y porcentaje parcial, sin respuestas."),
    ("Finalizar intento", "POST /v1/attempts/{id}/complete", "Tiempo, porcentaje, estrellas y versión."),
    ("Consultar progreso", "GET /v1/progress/me", "Constelaciones, actividades y desbloqueos."),
], [2200, 2900, 4260])
h2(doc, "Reglas transversales")
bullet(doc, "Versionar rutas y esquemas desde el primer contrato público.")
bullet(doc, "Validar un identificador de intento estable para procesar reenvíos de forma idempotente.")
bullet(doc, "Responder errores con código, mensaje seguro, correlationId y condición de reintento.")
bullet(doc, "No aceptar desde el cliente campos que el servidor pueda derivar, como identidad o institución.")
bullet(doc, "Registrar versión del juego y versión de la actividad con cada resultado.")
callout(doc, "RIESGO", "Porcentaje calculado en el cliente", "El requisito actual indica que el juego calcula el porcentaje. El backend debe validar rangos, consistencia y reglas básicas para reducir manipulación o datos imposibles.")

doc.add_page_break()

h1(doc, "6. Modelo de datos y propiedad")
callout(doc, "RECOMENDADO", "Esquema dedicado", "Usar un esquema lógico propio —por ejemplo caza_estelar— dentro de la base compartida. La web obtiene información mediante vistas o consultas acordadas; ningún sistema modifica directamente las tablas internas del otro.")
matrix(doc, ["Entidad", "Campos esenciales"], [
    ("game_version", "id, versionName, minimumSupported, publishedAt, status"),
    ("activity_version", "id, activityId, gameVersionId, grade, topic, status"),
    ("attempt", "attemptId, studentId, activityVersionId, number, status, startedAt, completedAt"),
    ("attempt_result", "attemptId, elapsedSeconds, percentage, stars, receivedAt"),
    ("progress", "studentId, activityId, bestAttemptId, bestPercentage, updatedAt"),
    ("constellation_progress", "studentId, constellationId, recoveredAt, starsRecovered"),
    ("audit_event", "eventId, actorId, action, targetRef, occurredAt, anonymizedAt"),
], [2700, 6660])
h2(doc, "Consistencia")
bullet(doc, "Clave única sobre attemptId para impedir duplicados.")
bullet(doc, "Transacción única para cerrar el intento, actualizar mejor resultado y recalcular progreso.")
bullet(doc, "El mejor resultado conserva el tiempo y número del mismo intento.")
bullet(doc, "Los alumnos sin resultado aportan 0 % al promedio, según la regla de reporte.")
bullet(doc, "Eliminar una cuenta exige retirar sus resultados y recalcular agregados.")
callout(doc, "PENDIENTE", "Motor relacional", "PostgreSQL administrado mediante Amazon RDS o Aurora es una opción natural, pero debe respetarse el motor ya elegido por la web y la base compartida.")

doc.add_page_break()

h1(doc, "7. Autenticación, autorización y privacidad")
h2(doc, "Identidad")
body(doc, "Amazon Cognito User Pools es una alternativa recomendada para evaluar: admite usuarios creados administrativamente, autenticación con usuario y contraseña, emisión de JWT y renovación mediante refresh tokens. Si la web ya dispone de un proveedor de identidad, el backend debe integrarse con ese proveedor en lugar de duplicar cuentas.")
h2(doc, "Autorización")
matrix(doc, ["Rol", "Permiso backend"], [
    ("Alumno", "Su sesión, progreso e intentos; nunca datos de otros alumnos."),
    ("Docente", "Resultados de secciones asignadas."),
    ("Director", "Resultados y exportaciones de su colegio."),
    ("Administrador general", "Usuarios, bloqueos, auditoría, traslados y eliminación."),
], [2700, 6660])
h2(doc, "Controles mínimos")
bullet(doc, "Cifrado TLS en tránsito y cifrado de la base, backups y snapshots en reposo.")
bullet(doc, "Secretos fuera del código, almacenados en AWS Secrets Manager y sujetos a rotación.")
bullet(doc, "Principio de mínimo privilegio para aplicación, migraciones, reportes y soporte.")
bullet(doc, "Logs sin contraseñas, tokens, fecha de nacimiento ni información innecesaria de menores.")
bullet(doc, "Auditoría anonimizada cuando se elimina definitivamente una cuenta.")
callout(doc, "PENDIENTE", "Cambio de contraseña", "Todavía no se ha decidido si las credenciales generadas deben cambiarse durante el primer ingreso.")

doc.add_page_break()

h1(doc, "8. Confiabilidad y operación")
h2(doc, "Idempotencia y reintentos")
body(doc, "El backend debe procesar varias veces la misma solicitud sin repetir sus efectos. attemptId funcionará como clave idempotente y el servidor devolverá el resultado previamente confirmado cuando reciba un reenvío equivalente.")
bullet(doc, "Rechazar la reutilización del mismo attemptId con contenido incompatible.")
bullet(doc, "Distinguir estados RECEIVED, PROCESSING, COMPLETED y REJECTED.")
bullet(doc, "Mantener correlationId desde la aplicación hasta logs y métricas.")
h2(doc, "Base de datos")
bullet(doc, "Backups automáticos y recuperación a un punto en el tiempo.")
bullet(doc, "Pruebas periódicas de restauración, no solo creación de snapshots.")
bullet(doc, "Evaluar Multi-AZ para producción según criticidad, costo y disponibilidad requerida.")
bullet(doc, "RDS Proxy si el runtime genera conexiones breves o picos impredecibles.")
h2(doc, "Observabilidad")
bullet(doc, "CloudWatch Logs con JSON estructurado y periodos de retención definidos.")
bullet(doc, "Métricas de latencia, errores, intentos rechazados, duplicados y cola de reintentos.")
bullet(doc, "Alarmas por tasa de error, caída de autenticación, agotamiento de conexiones y retraso de resultados.")
bullet(doc, "Objetivos operativos propuestos: disponibilidad, latencia p95 y éxito de procesamiento.")

doc.add_page_break()

h1(doc, "9. Ambientes, despliegue y pruebas")
matrix(doc, ["Ambiente", "Propósito", "Datos"], [
    ("Desarrollo", "Integración rápida y pruebas locales/remotas.", "Sintéticos; nunca copia libre de producción."),
    ("Pruebas", "QA, validación académica e integración completa.", "Instituciones y alumnos ficticios."),
    ("Producción", "Uso real de colegios y escuelas.", "Datos protegidos y auditados."),
], [2100, 3900, 3360])
h2(doc, "Infraestructura como código")
body(doc, "Se recomienda AWS CDK o una alternativa equivalente para definir redes, cómputo, secretos, permisos, observabilidad y parámetros por ambiente. AWS CDK provisiona recursos mediante CloudFormation y permite versionar la infraestructura junto con el código.")
h2(doc, "Pipeline sugerido")
number(doc, "Compilar, ejecutar análisis estático y pruebas unitarias.")
number(doc, "Ejecutar pruebas de contrato, integración e idempotencia.")
number(doc, "Desplegar automáticamente en desarrollo.")
number(doc, "Promover a pruebas y ejecutar QA funcional, bilingüe y de interrupciones.")
number(doc, "Solicitar aprobación de dirección del proyecto.")
number(doc, "Promover el mismo artefacto validado a producción.")
callout(doc, "RECOMENDADO", "Migraciones", "Usar migraciones versionadas, reversibles cuando sea posible y ejecutadas por una identidad separada del runtime de la aplicación.")

doc.add_page_break()

h1(doc, "10. Decisiones por resolver")
matrix(doc, ["Decisión", "Opciones", "Criterio"], [
    ("Framework Kotlin", "Ktor / Spring Boot", "Experiencia, arranque, observabilidad y productividad."),
    ("Runtime AWS", "ECS Fargate / Lambda", "Patrón de tráfico, costo, latencia y conexiones DB."),
    ("Identidad", "Cognito / proveedor existente", "Fuente única de usuarios y control institucional."),
    ("Motor de datos", "Motor existente / RDS PostgreSQL / Aurora", "Compatibilidad con la base compartida."),
    ("Contrato", "REST JSON / alternativa acordada", "Compatibilidad con Android y web."),
    ("Reportes", "Vistas directas / servicio / eventos", "Acoplamiento y actualización de la bitácora."),
    ("Credenciales", "Cambio inicial / sin cambio", "Operación escolar y seguridad."),
], [2400, 3300, 3660])
h2(doc, "Prueba de concepto recomendada")
bullet(doc, "Autenticación de una cuenta de alumno y renovación silenciosa de sesión.")
bullet(doc, "Creación y finalización idempotente de un intento.")
bullet(doc, "Actualización del mejor resultado y progreso de una constelación.")
bullet(doc, "Consulta de la bitácora sobre la misma base sin acceso cruzado de escritura.")
bullet(doc, "Despliegue reproducible en desarrollo y pruebas con métricas visibles.")
callout(doc, "RECOMENDADO", "Criterio de salida", "No elegir framework ni runtime por preferencia. Elegir la opción que cumpla el flujo completo con menor complejidad operativa, rendimiento suficiente y pruebas automatizadas claras.")

doc.add_page_break()

h1(doc, "11. Ruta de implementación")
number(doc, "Acordar límites de propiedad con el equipo de la web/bitácora y documentar el esquema compartido.")
number(doc, "Definir identidad, contrato HTTPS/JSON, tokens, roles y política de sesiones.")
number(doc, "Construir la prueba de concepto Kotlin/JVM y comparar Fargate frente a Lambda.")
number(doc, "Implementar vertical mínima: versión, autenticación, intento idempotente, resultado y progreso.")
number(doc, "Añadir auditoría, bloqueo remoto, traslados y eliminación de cuentas.")
number(doc, "Incorporar backups, restauración, secretos, métricas y alarmas.")
number(doc, "Ejecutar pruebas de carga, seguridad, contrato y recuperación antes de producción.")
h2(doc, "Recomendación final")
body(doc, "Iniciar con un backend Kotlin/JVM modular y un esquema dedicado dentro de la base compartida. Mantener una sola unidad de despliegue mientras el dominio se estabiliza. Evaluar ECS Fargate como opción base por su ajuste natural a un servicio persistente con base relacional, y conservar Lambda como alternativa si la prueba demuestra una carga intermitente, latencia aceptable y administración segura de conexiones.")
callout(doc, "CONFIRMADO", "Gobernanza", "Dirección del proyecto aprueba la publicación; desarrollo define y mantiene el contrato de integración con la API.")

doc.add_page_break()

h1(doc, "12. Fuentes técnicas consultadas")
body(doc, "Las recomendaciones de AWS se basan en documentación oficial consultada en agosto de 2026:")
sources = [
    ("Amazon Cognito User Pools", "https://docs.aws.amazon.com/cognito/latest/developerguide/cognito-user-pools.html"),
    ("Autenticación con Cognito", "https://docs.aws.amazon.com/cognito/latest/developerguide/authentication.html"),
    ("JWT authorizers en API Gateway", "https://docs.aws.amazon.com/apigateway/latest/developerguide/http-api-jwt-authorizer.html"),
    ("Cifrado de Amazon RDS", "https://docs.aws.amazon.com/AmazonRDS/latest/UserGuide/Overview.Encryption.html"),
    ("Backups automáticos de Amazon RDS", "https://docs.aws.amazon.com/AmazonRDS/latest/UserGuide/USER_WorkingWithAutomatedBackups.html"),
    ("Amazon RDS Proxy", "https://docs.aws.amazon.com/AmazonRDS/latest/UserGuide/rds-proxy.html"),
    ("Rotación en AWS Secrets Manager", "https://docs.aws.amazon.com/secretsmanager/latest/userguide/rotating-secrets.html"),
    ("CloudWatch Application Signals", "https://docs.aws.amazon.com/AmazonCloudWatch/latest/monitoring/CloudWatch-Application-Monitoring-Sections.html"),
    ("AWS CDK", "https://docs.aws.amazon.com/cdk/v2/guide/home.html"),
    ("AWS Lambda con Java", "https://docs.aws.amazon.com/lambda/latest/dg/lambda-java.html"),
    ("Decisión AWS Fargate o Lambda", "https://docs.aws.amazon.com/pdfs/decision-guides/latest/fargate-or-lambda/fargate-or-lambda.pdf"),
]
for title, url in sources:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r = p.add_run(title + ": ")
    set_font(r, 9.5, bold=True, color=INK)
    r = p.add_run(url)
    set_font(r, 9.2, color=BLUE)
body(doc, "Nota: los servicios concretos, el framework y el contrato final deberán aprobarse mediante la prueba de concepto y la coordinación con la web/bitácora. Esta propuesta no reemplaza una revisión legal o de seguridad especializada sobre datos de menores.")

doc.core_properties.title = "Propuesta de Backend - Caza Estelar"
doc.core_properties.subject = "Arquitectura, datos, seguridad e infraestructura AWS"
doc.core_properties.author = "Explorax"
doc.core_properties.keywords = "Caza Estelar, backend, AWS, Kotlin, API, RDS, Cognito"
doc.save(OUT)
print(OUT)
