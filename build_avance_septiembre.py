# -*- coding: utf-8 -*-
"""Genera el avance de septiembre 2026: planteamiento del juego y de la app administrativa.

Fuente reproducible de `Avance_Septiembre_Caza_Estelar.docx`.
Contenido derivado de:
  - caza-estelar/docs/DECISIONES_FASE_0.md
  - caza-estelar/docs/PLAN_SEPTIEMBRE_2026.md
  - caza-estelar/docs/PLAN_DESARROLLO.md
  - caza-estelar/docs/CONTENT_SCHEMA_V2.md
  - Avance_solicitado_Caza_Estelar.docx (Versión final planeada)

Requiere: python-docx y Pillow.  Uso:  python build_avance_septiembre.py
"""
from pathlib import Path
import math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Avance_Septiembre_Caza_Estelar.docx"
DIAG = ROOT / "qa_avance_septiembre" / "diagramas"
DIAG.mkdir(parents=True, exist_ok=True)

NAVY = "111B3C"; BLUE = "1967D2"; CYAN = "48CAE4"
INK = "202632"; MUTED = "5F6B7A"; PALE = "EEF4FF"
GREEN_BG = "E2F3EA"; AMBER_BG = "FFF4D6"


# --------------------------------------------------------------------- diagramas
def font(size=28, bold=False):
    name = "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"
    if Path(name).exists():
        return ImageFont.truetype(name, size)
    return ImageFont.load_default()


def diagram(name, title, nodes, edges, size=(1500, 820)):
    im = Image.new("RGB", size, "#F8FAFE")
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((22, 22, size[0] - 22, size[1] - 22), radius=26,
                        fill="#F8FAFE", outline="#CAD7EF", width=3)
    d.text((56, 40), title, fill="#111B3C", font=font(36, True))
    boxes = {}
    for key, x, y, w, h, label, color in nodes:
        boxes[key] = (x, y, w, h)
        d.rounded_rectangle((x, y, x + w, y + h), radius=16,
                            fill=color, outline="#315B8A", width=3)
        lines = label.split("\n")
        fs = 22 if len(lines) <= 3 else 19
        total = sum(d.textbbox((0, 0), t, font=font(fs, True))[3] for t in lines) + 6 * (len(lines) - 1)
        yy = y + (h - total) / 2
        for t in lines:
            bb = d.textbbox((0, 0), t, font=font(fs, True))
            d.text((x + (w - (bb[2] - bb[0])) / 2, yy), t, fill="#172033", font=font(fs, True))
            yy += bb[3] - bb[1] + 6
    for a, b, label in edges:
        ax, ay, aw, ah = boxes[a]
        bx, by, bw, bh = boxes[b]
        if by > ay + ah:
            p1 = (ax + aw / 2, ay + ah); p2 = (bx + bw / 2, by)
        elif by + bh < ay:
            p1 = (ax + aw / 2, ay); p2 = (bx + bw / 2, by + bh)
        elif bx > ax:
            p1 = (ax + aw, ay + ah / 2); p2 = (bx, by + bh / 2)
        else:
            p1 = (ax, ay + ah / 2); p2 = (bx + bw, by + bh / 2)
        d.line((p1, p2), fill="#485A72", width=4)
        ang = math.atan2(p2[1] - p1[1], p2[0] - p1[0]); L = 15
        d.polygon([p2,
                   (p2[0] - L * math.cos(ang - .5), p2[1] - L * math.sin(ang - .5)),
                   (p2[0] - L * math.cos(ang + .5), p2[1] - L * math.sin(ang + .5))],
                  fill="#485A72")
        if label:
            mx = (p1[0] + p2[0]) / 2; my = (p1[1] + p2[1]) / 2 - 20
            bb = d.textbbox((0, 0), label, font=font(17))
            d.rectangle((mx - 6, my - 2, mx + bb[2] + 6, my + bb[3] + 2), fill="#F8FAFE")
            d.text((mx, my), label, fill="#485A72", font=font(17))
    path = DIAG / f"{name}.png"
    im.save(path)
    return path


flow = diagram("flujo", "Flujo principal del juego", [
    ("load", 60, 150, 210, 92, "Carga de\nrecursos", "#DDEBFF"),
    ("login", 340, 150, 210, 92, "Inicio de\nsesión (API)", "#DDEBFF"),
    ("menu", 620, 150, 210, 92, "Menú\nprincipal", "#DDEBFF"),
    ("map", 900, 150, 210, 92, "Mapa de\nconstelaciones", "#DDEBFF"),
    ("levels", 1180, 150, 240, 92, "Lista de\nniveles", "#DDEBFF"),
    ("act", 1180, 400, 240, 92, "Actividad\ninteractiva", "#FFF2CB"),
    ("result", 900, 400, 210, 92, "Resultado y\nestrellas", "#DFF4E9"),
    ("save", 620, 400, 210, 92, "Enviar progreso\na la API", "#DFF4E9"),
], [
    ("load", "login", ""), ("login", "menu", "credenciales"), ("menu", "map", "jugar"),
    ("map", "levels", "constelación"), ("levels", "act", "elegir nivel"),
    ("act", "result", "finalizar"), ("result", "save", "score / tiempo"),
    ("save", "levels", "otro nivel"),
], size=(1500, 560))

hier = diagram("jerarquia", "Jerarquía de contenido v2", [
    ("cat", 560, 60, 360, 90, "Catálogo\nschemaVersion 2, revision", "#EAE3FF"),
    ("c1", 120, 250, 340, 100, "Constelación / tema\nid, version, area, title", "#DDEBFF"),
    ("c2", 1020, 250, 340, 100, "Constelación / tema\n(otra área)", "#DDEBFF"),
    ("l1", 60, 470, 240, 96, "Estrella / nivel 1\nmechanic, completion", "#FFF2CB"),
    ("l2", 330, 470, 240, 96, "Estrella / nivel 2\nunlock.prerequisiteLevelIds", "#FFF2CB"),
    ("l3", 600, 470, 240, 96, "Estrella / nivel N", "#FFF2CB"),
], [
    ("cat", "c1", ""), ("cat", "c2", ""),
    ("c1", "l1", ""), ("c1", "l2", ""), ("c1", "l3", ""),
], size=(1440, 620))

entities = diagram("entidades", "UML de entidades (juego y administración)", [
    ("user", 60, 130, 250, 110, "Usuario\nid, rol, institutionId", "#DDEBFF"),
    ("student", 420, 130, 250, 110, "Alumno\nfechaNac, grado, sección", "#DDEBFF"),
    ("inst", 780, 130, 250, 110, "Institución\npaís, secciones", "#DDEBFF"),
    ("const", 1140, 130, 250, 110, "Constelación\narea, version", "#EAE3FF"),
    ("level", 1140, 380, 250, 110, "Nivel\nmechanic, completion", "#EAE3FF"),
    ("attempt", 780, 380, 250, 110, "Intento\nscore, durationMs", "#DFF4E9"),
    ("progress", 420, 380, 250, 110, "Progreso\nbestScore, stars, completed", "#DFF4E9"),
    ("asset", 60, 380, 250, 110, "RecursoMultimedia\nmimeType, sha256", "#EAE3FF"),
], [
    ("user", "student", "perfil"), ("student", "inst", "pertenece"),
    ("student", "progress", "obtiene"), ("const", "level", "contiene"),
    ("level", "progress", "genera"), ("progress", "attempt", "1..n"),
    ("level", "asset", "usa"),
], size=(1460, 540))

variables = diagram("variables", "UML de variables y estado del juego", [
    ("session", 60, 130, 320, 150,
     "PlayerSession\nuserId, token\nauthenticated, role\ninstitutionId", "#DDEBFF"),
    ("settings", 430, 130, 320, 150,
     "AppSettings\nlanguage\nsoundEnabled\nmusicVolume, effectsVolume", "#DDEBFF"),
    ("game", 800, 130, 320, 150,
     "ActivitySession\ntargets/opciones\nscore, elapsedMs\ncompleted", "#FFF2CB"),
    ("prog", 430, 400, 320, 160,
     "LevelProgress\nlevelId, levelVersion\nbestScore, attemptsUsed\ndurationMs, completed\nstars (derivado)", "#DFF4E9"),
], [
    ("session", "settings", "usa"), ("settings", "game", "configura"),
    ("game", "prog", "actualiza"), ("session", "prog", "asocia"),
], size=(1200, 620))

usecases = diagram("casos_uso", "Casos de uso", [
    ("student", 40, 210, 190, 100, "Alumno", "#DDEBFF"),
    ("uc_login", 290, 90, 360, 74, "Iniciar sesión en línea", "#EEF4FF"),
    ("uc_play", 290, 200, 360, 74, "Elegir constelación y nivel; jugar", "#EEF4FF"),
    ("uc_prog", 290, 310, 360, 74, "Consultar su progreso", "#EEF4FF"),
    ("author", 40, 560, 190, 100, "Autor de contenido", "#FFF2CB"),
    ("uc_edit", 290, 520, 360, 74, "Crear y editar constelaciones y niveles", "#FDF5DF"),
    ("reviewer", 40, 670, 190, 100, "Revisor académico", "#DFF4E9"),
    ("uc_review", 290, 630, 360, 74, "Aprobar o solicitar cambios", "#EAF8F0"),
    ("admin", 760, 210, 190, 100, "Administrador general", "#EAE3FF"),
    ("uc_users", 1010, 120, 380, 74, "Gestionar usuarios, instituciones y bloqueos", "#F2EDFF"),
    ("uc_pub", 1010, 230, 380, 74, "Publicar revisión inmutable del catálogo", "#F2EDFF"),
    ("uc_rep", 1010, 340, 380, 74, "Reportes globales PDF / Excel", "#F2EDFF"),
], [
    ("student", "uc_login", ""), ("student", "uc_play", ""), ("student", "uc_prog", ""),
    ("author", "uc_edit", ""), ("reviewer", "uc_review", ""),
    ("admin", "uc_users", ""), ("admin", "uc_pub", ""), ("admin", "uc_rep", ""),
], size=(1460, 820))

adminarch = diagram("arquitectura_admin", "Arquitectura de la aplicación administrativa", [
    ("browser", 60, 120, 250, 100, "Navegador\nadministrativo", "#DDEBFF"),
    ("app", 60, 400, 250, 100, "App Android\n(juego)", "#DDEBFF"),
    ("nginx", 470, 250, 240, 100, "Nginx\nHTTPS, proxy", "#FFF2CB"),
    ("web", 850, 110, 260, 100, "Frontend React + TS\n(archivos estáticos)", "#EAE3FF"),
    ("api", 850, 260, 260, 100, "API NestJS\n(Node.js, Docker)", "#EAE3FF"),
    ("db", 850, 430, 260, 100, "PostgreSQL\n+ Prisma", "#DFF4E9"),
    ("files", 1230, 340, 220, 100, "Volumen de\narchivos", "#DFF4E9"),
], [
    ("browser", "nginx", "HTTPS"), ("app", "nginx", "HTTPS /api"),
    ("nginx", "web", ""), ("nginx", "api", "/api"),
    ("api", "db", ""), ("api", "files", ""),
], size=(1500, 600))


# --------------------------------------------------------------------- docx
def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), color); tcPr.append(shd)


def margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc.get_or_add_tcPr()
    m = tc.first_child_found_in("w:tcMar")
    if m is None:
        m = OxmlElement("w:tcMar"); tc.append(m)
    for tag, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        e = OxmlElement("w:" + tag); e.set(qn("w:w"), str(val)); e.set(qn("w:type"), "dxa"); m.append(e)


def repeat(row):
    trPr = row._tr.get_or_add_trPr()
    x = OxmlElement("w:tblHeader"); x.set(qn("w:val"), "true"); trPr.append(x)


def no_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = Inches(.75); sec.bottom_margin = Inches(.7)
sec.left_margin = Inches(.85); sec.right_margin = Inches(.85)

styles = doc.styles
for s in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    st = styles[s]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor.from_string(INK)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for s, size, color, before, after in [
    ("Title", 27, NAVY, 0, 6), ("Subtitle", 13, MUTED, 0, 12),
    ("Heading 1", 17, NAVY, 15, 6), ("Heading 2", 12.5, BLUE, 10, 4),
    ("Heading 3", 11, INK, 7, 3),
]:
    st = styles[s]
    st.font.size = Pt(size); st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = s != "Subtitle"
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]
header.text = "EXPLORAX  /  CAZA ESTELAR"
header.runs[0].font.size = Pt(8); header.runs[0].font.bold = True
header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("Avance de septiembre 2026  |  ")
fr.font.size = Pt(8); fr.font.color.rgb = RGBColor.from_string(MUTED)
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); footer._p.append(fld)


def kicker(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(42)
    r = p.add_run(text); r.bold = True; r.font.name = "Arial"
    r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(CYAN)


def para(text, lead=None, lead_color=BLUE):
    p = doc.add_paragraph()
    if lead:
        r = p.add_run(lead + " "); r.bold = True
        r.font.color.rgb = RGBColor.from_string(lead_color)
    p.add_run(text)
    return p


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def figure(path, alt):
    doc.add_picture(str(path), width=Inches(6.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = doc.inline_shapes[-1]
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", alt.split(".")[0])


def callout(text, fill=AMBER_BG):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.LEFT
    shade(t.cell(0, 0), fill); margins(t.cell(0, 0), 150, 170, 150, 170)
    t.cell(0, 0).text = text


def table(headers, rows, widths, shade_col=None, shade_map=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, (h, w) in enumerate(zip(headers, widths)):
        t.columns[i].width = Inches(w)
        c = t.cell(0, i); c.text = h; shade(c, NAVY)
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    repeat(t.rows[0]); no_split(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        no_split(t.rows[-1])
        for i, v in enumerate(row):
            cells[i].text = v
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cells[i])
        if shade_col is not None and shade_map:
            shade(cells[shade_col], shade_map.get(row[shade_col], "FFFFFF"))
    return t


# --------------------------------------------------------------------- portada
kicker("AVANCE DE SEPTIEMBRE 2026")
doc.add_paragraph("Caza Estelar", style="Title")
doc.add_paragraph("Planteamiento del juego y de la aplicación administrativa", style="Subtitle")

para("septiembre no implementa la aplicación. El avance del mes es este planteamiento de "
     "diseño más una prueba mínima de concepto verificable del contrato de contenido v2. "
     "La implementación no forma parte de este avance.", lead="Alcance del mes:", lead_color=NAVY)

t = doc.add_table(rows=5, cols=2); t.alignment = WD_TABLE_ALIGNMENT.LEFT; t.autofit = False
for i, (a, b) in enumerate([
    ("Proyecto", "Caza Estelar — juego educativo Android y aplicación administrativa"),
    ("Periodo", "1 al 30 de septiembre de 2026"),
    ("Entregable del mes", "Documento de planteamiento + prueba de concepto del contrato v2"),
    ("Estado", "Fases 0 y 1 completadas; planteamiento de Fase 2 en este documento"),
    ("Corte", "8 de septiembre de 2026"),
]):
    t.columns[0].width = Inches(1.7); t.columns[1].width = Inches(4.7)
    t.cell(i, 0).text = a; t.cell(i, 1).text = b; shade(t.cell(i, 0), PALE)
    for c in t.rows[i].cells:
        margins(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    t.cell(i, 0).paragraphs[0].runs[0].bold = True

doc.add_paragraph("Nota de alcance", style="Heading 2")
para("Este documento distingue lo ya ejecutado de lo planteado. IMPLEMENTADO = existe y "
     "compila en el repositorio. PLANTEADO = diseño acordado, aún no implementado. "
     "PROTOTIPO = demostración parcial.")
callout("Prueba mínima de concepto incluida este mes: modelos Kotlin del contrato v2 "
        "(ContentModelsV2.kt), esquema JSON (catalog-v2.schema.json) y sus pruebas. Compilan "
        "y pasan, pero no están conectados al parser, la navegación ni el progreso.")
doc.add_page_break()

# --------------------------------------------------------------------- 1
doc.add_paragraph("1. Resumen del avance", style="Heading 1")
para("Durante esta etapa se cerraron las decisiones de producto pendientes (Fase 0) y se "
     "estableció un entorno de trabajo reproducible (Fase 1): repositorio Git, Gradle Wrapper, "
     "toolchain embebido y una línea base de pruebas verde. Sobre esa base, el avance del mes "
     "es el planteamiento del contrato de contenido v2, del modelo de progreso v2, de la "
     "arquitectura de la aplicación administrativa y de la estructura de repositorios, más una "
     "prueba de concepto del modelo v2 en código.")
para("La implementación del parser v2, la navegación por niveles, el retiro del modo sin "
     "conexión y el backend quedan fuera del alcance de septiembre.")

# --------------------------------------------------------------------- 2
doc.add_paragraph("2. Trabajo de base ya ejecutado", style="Heading 1")
doc.add_paragraph("2.1 Decisiones de producto (Fase 0)", style="Heading 2")
table(["#", "Tema", "Decisión"], [
    ("1", "Contrato de contenido", "Migrar directo a v2 (constelación → nivel), sin capa de compatibilidad v1."),
    ("2", "Motor / arquitectura móvil", "Sin decidir. Spike comparativo y ADR antes de fijar. Sin reescritura total aprobada."),
    ("3", "Conectividad", "Solo en línea. Se retira la cola de sincronización persistente y el juego sin conexión. El catálogo sigue empaquetado en la APK."),
    ("4", "Repositorios", "Repo A = app Android. Repo B = aplicación administrativa (NestJS + React + Prisma + infra, juntos)."),
    ("5", "Modelo de progreso", "Campos v2: levelId, levelVersion, completed, score, attemptsUsed, durationMs, updatedAt. Porcentaje y estrellas derivados de score."),
    ("6", "Narración de voz", "narrationAssetId se mantiene opcional en el contrato v2."),
    ("7", "Reproducibilidad", "Fase 1 ejecutada el 7 de septiembre."),
    ("8", "Stack administrativo", "React + TypeScript, Node.js + NestJS, PostgreSQL + Prisma, Docker Compose, Nginx, JWT + Argon2, GitHub Actions, VPS Ubuntu, dominio explorax.app, ~20 000 usuarios. Reemplaza 'backend Kotlin' y 'AWS' de documentos anteriores."),
], [0.4, 1.7, 4.3])

doc.add_paragraph("2.2 Entorno reproducible (Fase 1)", style="Heading 2")
table(["Elemento", "Estado"], [
    ("Repositorio Git en la raíz, rama main, con .gitignore", "IMPLEMENTADO"),
    ("Carpeta del proyecto renombrada a caza-estelar/", "IMPLEMENTADO"),
    ("Gradle Wrapper fijado a Gradle 8.9 (compatible con AGP 8.7.3)", "IMPLEMENTADO"),
    ("Toolchain embebido: JDK 17, Android SDK 35, adb", "IMPLEMENTADO"),
    ("Línea base: core:test 38/38 verde; assembleDebug correcto", "IMPLEMENTADO"),
    ("verifyProductionConfiguration bloquea release con configuración de desarrollo", "IMPLEMENTADO"),
], [4.6, 1.8], shade_col=1, shade_map={"IMPLEMENTADO": GREEN_BG})

doc.add_page_break()

# --------------------------------------------------------------------- 3
doc.add_paragraph("3. Planteamiento de la aplicación del juego", style="Heading 1")

doc.add_paragraph("3.1 Flujo principal", style="Heading 2")
figure(flow, "Flujo del juego: carga, inicio de sesión en línea, menú, mapa de constelaciones, lista de niveles, actividad, resultado y envío de progreso a la API.")
para("El juego exige conexión para iniciar sesión y entrar. El contenido educativo viaja "
     "empaquetado en la APK; los resultados se envían a la API al terminar cada nivel.")

doc.add_paragraph("3.2 Jerarquía de contenido v2", style="Heading 2")
figure(hier, "Jerarquía v2: el catálogo contiene constelaciones (temas matemáticos) y cada constelación contiene estrellas o niveles jugables. No hay capa intermedia de temas.")
para("Una constelación representa un tema matemático; una estrella representa un nivel "
     "jugable. Al aprobar un nivel se conquista su estrella. Se elimina la capa de temas del "
     "contrato v1.")

doc.add_paragraph("3.3 UML de entidades", style="Heading 2")
figure(entities, "Entidades: Usuario, Alumno, Institución, Constelación, Nivel, Intento, Progreso y RecursoMultimedia.")
para("Usuario es la cuenta de acceso común; Alumno es el perfil académico vinculado a una "
     "cuenta de alumno. Constelación y Nivel provienen del contrato de contenido; Progreso e "
     "Intento registran el desempeño (Progreso conserva el mejor resultado; Intento conserva "
     "el historial).")

doc.add_paragraph("3.4 UML de variables y estado", style="Heading 2")
figure(variables, "Variables de PlayerSession, AppSettings, ActivitySession y LevelProgress con los campos del modelo v2.")
para("LevelProgress guarda score, attemptsUsed, durationMs y completed. El porcentaje y las "
     "estrellas se derivan de score y no se almacenan por separado en el cliente.")

doc.add_paragraph("3.5 Casos de uso", style="Heading 2")
figure(usecases, "Casos de uso para Alumno, Autor de contenido, Revisor académico y Administrador general.")

doc.add_paragraph("3.6 Programación de eventos", style="Heading 2")
table(["Evento", "Respuesta planteada"], [
    ("Inicio de la aplicación", "Precargar recursos y comprobar sesión guardada; exigir conexión."),
    ("Inicio de sesión", "Enviar credenciales a la API por HTTPS y crear la sesión según el rol."),
    ("Selección de nivel", "Consultar desbloqueos (prerequisiteLevelIds), cargar contenido y comenzar el temporizador."),
    ("Respuesta del alumno", "Evaluar, actualizar score y mostrar retroalimentación de Tojil (sin marcar error directo)."),
    ("Pérdida de conexión", "Pausar la actividad y el temporizador; mostrar aviso hasta recuperar la conexión."),
    ("Fin del nivel", "Calcular score, derivar porcentaje y estrellas, guardar localmente y enviar el evento idempotente a la API."),
    ("Bloqueo remoto de la cuenta", "Cerrar la sesión de inmediato y eliminar intentos incompletos locales."),
    ("Cambio de idioma o sonido", "Guardar la preferencia y actualizar la interfaz."),
], [1.9, 4.5])

doc.add_paragraph("3.7 Distribución de pantallas", style="Heading 2")
table(["Pantalla", "Distribución planteada"], [
    ("Carga", "Logo y animación al centro; progreso abajo; fondo espacial."),
    ("Acceso", "Escena de constelaciones al fondo; formulario en panel lateral; ajustes en la esquina superior izquierda."),
    ("Menú", "Observatorio al fondo; botones principales a la derecha."),
    ("Mapa de constelaciones", "Constelaciones en el área central; estado de bloqueo y progreso junto a cada una."),
    ("Lista de niveles", "Niveles de la constelación en desplazamiento vertical; estados disponible, bloqueado y completado."),
    ("Actividad", "Interacción al centro; instrucción arriba; tiempo y pausa en los bordes."),
    ("Resultado", "Porcentaje y estrellas al centro; botones para repetir, continuar o volver al mapa."),
], [1.5, 4.9])

doc.add_paragraph("3.8 Mecánicas", style="Heading 2")
table(["Mecánica", "Estado", "Descripción"], [
    ("telescope_hunt", "IMPLEMENTADA", "Orientar la mira y capturar objetivos. Config: targetCount, lockRadius."),
    ("multiple_choice", "IMPLEMENTADA", "Selección de una respuesta correcta entre 2 y 6 opciones bilingües."),
    ("conteo de objetos", "PLANTEADA", "Contar elementos y elegir o escribir la cantidad."),
    ("ordenamiento de números", "PLANTEADA", "Ordenar una secuencia de números."),
    ("resolución de operaciones", "PLANTEADA", "Resolver una operación simple."),
    ("arrastre / emparejamiento", "PLANTEADA", "Arrastrar elementos a su lugar o emparejar pares."),
], [1.9, 1.3, 3.2], shade_col=1,
   shade_map={"IMPLEMENTADA": GREEN_BG, "PLANTEADA": PALE})

doc.add_page_break()

# --------------------------------------------------------------------- 4
doc.add_paragraph("4. Contrato de contenido v2 (planteamiento)", style="Heading 1")
para("Especificación completa en caza-estelar/docs/CONTENT_SCHEMA_V2.md. Reemplaza el "
     "contrato v1. Cambios principales:")
for x in [
    "constellation → levels[] (se elimina themes[]).",
    "MissionDefinition pasa a LevelDefinition; mission.version pasa a level.version.",
    "rewards.maximumStars se sustituye por completion (minimum_score) y reward (experience, badgeId).",
    "unlock.prerequisiteMissionIds pasa a unlock.prerequisiteLevelIds.",
    "Se añaden constellation.area, constellation.visual y minimumAppVersion.",
    "narrationAssetId se conserva como campo opcional.",
]:
    bullet(x)

doc.add_paragraph("4.1 Reglas de validación en el juego", style="Heading 2")
for x in [
    "schemaVersion soportado (2) y revision estrictamente mayor que la activa.",
    "Respuesta HTTPS de tamaño máximo 1 MiB.",
    "Cada mechanic.type registrado en la app; configuración tipada válida.",
    "prerequisiteLevelIds existentes, sin duplicados, sin ciclos y sin auto-referencia.",
    "Al menos un nivel por constelación publicada.",
    "Cada referencia de media existe en assets y es del tipo correcto.",
    "Ante cualquier fallo, se conserva el último catálogo válido.",
]:
    bullet(x)

doc.add_paragraph("4.2 Modelo de progreso v2", style="Heading 2")
table(["Campo", "Descripción"], [
    ("levelId, levelVersion", "Identidad del nivel y versión con la que se jugó."),
    ("completed", "El nivel fue aprobado (score suficiente)."),
    ("score", "Puntuación 0..100 del mejor intento."),
    ("attemptsUsed", "Número de intentos del mejor resultado."),
    ("durationMs", "Duración del mejor intento en milisegundos."),
    ("updatedAt", "Marca de tiempo de la última actualización."),
    ("percentage, stars", "Derivados de score. Rangos: 0–29 → 0★, 30–59 → 1★, 60–89 → 2★ y aprobado, 90–100 → 3★."),
], [1.8, 4.6])
para("El evento de sincronización lleva un eventId (UUID) para que el servidor lo procese de "
     "forma idempotente. Contrato completo en REMOTE_API_CONTRACT.md.")

doc.add_page_break()

# --------------------------------------------------------------------- 5
doc.add_paragraph("5. Planteamiento de la aplicación administrativa", style="Heading 1")
para("Stack fijado por el Avance solicitado (Versión final planeada). Reemplaza la propuesta "
     "anterior de backend en Kotlin y de infraestructura AWS.")

doc.add_paragraph("5.1 Arquitectura", style="Heading 2")
figure(adminarch, "Arquitectura: navegador y app Android acceden por HTTPS a Nginx, que sirve el frontend React y enruta /api a la API NestJS, que usa PostgreSQL con Prisma y un volumen de archivos.")

doc.add_paragraph("5.2 Stack y herramientas", style="Heading 2")
table(["Herramienta", "Uso"], [
    ("React + TypeScript + Vite", "Interfaz administrativa web."),
    ("Node.js + NestJS", "API REST, reglas de negocio, autenticación y permisos."),
    ("PostgreSQL + Prisma", "Base de datos central y acceso tipado con migraciones."),
    ("Docker Compose", "Instalación reproducible de web, API y base de datos en el VPS."),
    ("Nginx + Let's Encrypt", "Proxy inverso, publicación del frontend y HTTPS."),
    ("JWT + Argon2", "Sesiones y almacenamiento protegido de contraseñas."),
    ("ExcelJS + PDFKit", "Carga de Excel/CSV y generación de reportes PDF/Excel."),
    ("GitHub Actions", "Pruebas y despliegue automatizado al VPS."),
], [2.1, 4.3])

doc.add_paragraph("5.3 Roles y permisos", style="Heading 2")
table(["Rol", "Permisos planteados"], [
    ("Alumno", "Acceder al juego y consultar su propio progreso."),
    ("Docente", "Consultar las secciones asignadas y generar reportes académicos."),
    ("Director", "Consultar todos los grados y secciones de su institución."),
    ("Administrador institucional", "Administrar alumnos, docentes y secciones de su institución."),
    ("Autor de contenido", "Crear y editar constelaciones, niveles y recursos."),
    ("Revisor académico", "Aprobar, rechazar y solicitar cambios al contenido."),
    ("Administrador general", "Administrar países, instituciones, usuarios, roles, bloqueos, contenido y auditoría global."),
], [1.9, 4.5])
para("Pendiente: matriz exacta de permisos y si autor y revisor deben ser personas distintas.")

doc.add_paragraph("5.4 Superficie de APIs", style="Heading 2")
table(["API / módulo", "Funciones planteadas"], [
    ("Autenticación", "Inicio y cierre de sesión, renovación de token, bloqueo y cierre remoto."),
    ("Usuarios e instituciones", "Altas individuales y masivas (Excel/CSV), grados, secciones y asignaciones."),
    ("Contenido (autoría)", "CRUD de constelaciones y niveles, catálogo de mecánicas, borradores, revisión y publicación."),
    ("Distribución", "Catálogo v2 publicado, de solo lectura, con ETag."),
    ("Multimedia", "Carga firmada y verificación de MIME, tamaño y hash."),
    ("Progreso", "Registro idempotente de intentos: score, tiempo y estrellas."),
    ("Reportes", "Filtros por fecha, grado, sección y actividad; exportación PDF/Excel."),
    ("Auditoría", "Registro de accesos, bloqueos, cambios y publicaciones."),
], [1.9, 4.5])

doc.add_paragraph("5.5 Esquema de base de datos", style="Heading 2")
table(["Grupo", "Tablas principales"], [
    ("Identidad", "users, roles, permissions, user_roles, sessions"),
    ("Instituciones", "countries, institutions, grades, sections, enrollments, teacher_assignments"),
    ("Contenido", "catalogs, constellations, activities, questions, answers, media_assets, content_versions"),
    ("Resultados", "attempts, activity_results, level_progress, constellation_progress"),
    ("Operación", "sync_events, account_blocks, audit_logs, report_jobs"),
], [1.5, 4.9])
para("La app nunca se conecta directamente a PostgreSQL: toda operación pasa por la API REST. "
     "El publicador transforma una revisión aprobada en el catálogo JSON v2 inmutable.")

doc.add_paragraph("5.6 Despliegue en VPS", style="Heading 2")
table(["Componente", "Instalación en el VPS"], [
    ("Nginx", "Publica el frontend y enruta /api hacia NestJS; termina HTTPS."),
    ("Frontend", "Archivos estáticos compilados desde React + TypeScript."),
    ("API", "Contenedor Docker con Node.js y NestJS."),
    ("Base de datos", "Contenedor PostgreSQL con volumen persistente y acceso interno."),
    ("Archivos", "Volumen protegido; migrable a almacenamiento de objetos más adelante."),
    ("Seguridad", "Firewall, SSH con llaves, HTTPS, variables secretas y actualizaciones."),
    ("Respaldo", "Copia diaria cifrada de PostgreSQL y archivos a una ubicación externa."),
], [1.5, 4.9])
para("Dominio: explorax.app. Dimensionado para ~20 000 usuarios registrados; la capacidad "
     "inicial se ajustará al confirmarse cuántos podrán conectarse simultáneamente.")

doc.add_page_break()

# --------------------------------------------------------------------- 6
doc.add_paragraph("6. Conectividad, seguridad y privacidad", style="Heading 1")
for x in [
    "Solo en línea: se exige conexión para iniciar sesión y entrar al juego; una desconexión durante una actividad la pausa.",
    "Todo el tráfico por HTTPS. El token de estudiante es de corta duración y en Android se cifra con Android Keystore (AES-GCM).",
    "Ningún token administrativo ni endpoint de escritura se incluye en la APK.",
    "Separación de APIs: autoría, distribución y progreso.",
    "Publicación separada de la edición; revisiones inmutables; registro de auditoría.",
    "Verificación de esquema, MIME, tamaño y hash de cada recurso en el servidor.",
    "Datos de menores: consentimiento gestionado por la institución; retención, eliminación y auditoría por definir con revisión legal (Guatemala y Estados Unidos).",
]:
    bullet(x)

# --------------------------------------------------------------------- 7
doc.add_paragraph("7. Estructura de repositorios", style="Heading 1")
table(["Repositorio", "Contenido", "Tecnología"], [
    ("Repo A — caza-estelar", "App Android del juego.", "Kotlin / Java / libGDX, Gradle."),
    ("Repo B — aplicación administrativa", "Backend, consola web e infraestructura, juntos y desplegables en el VPS.", "NestJS, React + TS, Prisma, Docker Compose."),
], [2.1, 2.9, 1.4])
para("El contrato de contenido (JSON Schema v2 y fixtures) se comparte entre ambos como "
     "paquete versionado.")

# --------------------------------------------------------------------- 8
doc.add_paragraph("8. Prueba mínima de concepto (anexo)", style="Heading 1")
para("Trabajo en código entregado este mes. Compila y sus pruebas pasan; no está conectado al "
     "parser, la navegación ni el progreso del juego.")
table(["Artefacto", "Ruta", "Contenido"], [
    ("Modelos v2", "core/.../content/v2/model/ContentModelsV2.kt",
     "ContentCatalogV2, ConstellationDefinitionV2, LevelDefinition, LevelCompletion, LevelReward, LevelUnlock, LevelMedia, LevelProgress."),
    ("Esquema JSON", "docs/contracts/catalog-v2.schema.json",
     "JSON Schema 2020-12 del catálogo v2, con validación condicional por mecánica."),
    ("Pruebas", "core/src/test/.../content/v2/ContentModelsV2Test.kt",
     "Resolución de niveles por id y derivación de estrellas/porcentaje por rangos."),
    ("Especificación", "docs/CONTENT_SCHEMA_V2.md",
     "Contrato v2 completo: estructura, ejemplo, mecánicas, validación y progreso."),
], [1.3, 2.6, 2.5])

# --------------------------------------------------------------------- 9
doc.add_paragraph("9. Aspectos por aclarar", style="Heading 1")
for x in [
    "Proveedor, país y capacidad inicial del VPS; usuarios simultáneos y volumen de multimedia.",
    "Propiedad de las tablas comunes con la bitácora web central y esquema compartido.",
    "Matriz definitiva de permisos por rol; si autor y revisor deben ser personas distintas.",
    "Formato exacto de los archivos Excel/CSV para carga masiva y diseño de los reportes PDF/Excel.",
    "Temas matemáticos, constelaciones y actividades: se plantearán con el equipo académico.",
    "Decisión de motor del juego (libGDX, Compose/Canvas o híbrido): spike y ADR.",
    "Procedimiento de publicación del APK: descarga privada, MDM institucional o Google Play.",
    "Política de conservación de historial, respaldos y datos personales de menores.",
]:
    bullet(x)

# --------------------------------------------------------------------- 10
doc.add_paragraph("10. Alcance de la implementación", style="Heading 1")
para("La implementación de este planteamiento —parser v2, validación, progreso v2, retiro del "
     "modo sin conexión, navegación por niveles, backend y consola— no forma parte del avance "
     "de septiembre. Su planificación se abordará una vez que el planteamiento y el contrato "
     "v2 estén aprobados. La hoja de ruta por fases, sin calendario, está en "
     "caza-estelar/docs/PLAN_DESARROLLO.md.")

doc.add_paragraph("Documentos de referencia", style="Heading 2")
for x in [
    "caza-estelar/docs/DECISIONES_FASE_0.md — decisiones de producto cerradas.",
    "caza-estelar/docs/CONTENT_SCHEMA_V2.md — contrato de contenido v2.",
    "caza-estelar/docs/PLAN_SEPTIEMBRE_2026.md — plan operativo del mes.",
    "caza-estelar/docs/PLAN_DESARROLLO.md — hoja de ruta por fases.",
    "Avance_solicitado_Caza_Estelar.docx — versión final planeada (fuente del stack y los repos).",
]:
    bullet(x)

for tbl in doc.tables:
    repeat(tbl.rows[0])
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.name = "Arial"
                    r._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
                    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
                    r.font.size = Pt(9)

doc.core_properties.title = "Avance de septiembre 2026 - Caza Estelar"
doc.core_properties.subject = "Planteamiento del juego y de la aplicación administrativa"
doc.core_properties.author = "Equipo Caza Estelar"
doc.save(OUT)
print(OUT)
