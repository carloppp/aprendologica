from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'Avance_solicitado_Caza_Estelar.docx'
IMG=ROOT/'qa_informe_avances'/'diagramas'
NAVY='111B3C'; BLUE='1967D2'; PALE='EEF4FF'; INK='202632'; MUTED='5F6B7A'

def shade(cell,color):
    pr=cell._tc.get_or_add_tcPr(); x=OxmlElement('w:shd'); x.set(qn('w:fill'),color); pr.append(x)
def margins(cell,v=100,h=120):
    pr=cell._tc.get_or_add_tcPr(); m=OxmlElement('w:tcMar')
    for tag,val in [('top',v),('start',h),('bottom',v),('end',h)]:
        e=OxmlElement('w:'+tag); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); m.append(e)
    pr.append(m)
def repeat(row):
    p=row._tr.get_or_add_trPr(); x=OxmlElement('w:tblHeader'); x.set(qn('w:val'),'true'); p.append(x)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,(h,w) in enumerate(zip(headers,widths)):
        t.columns[i].width=Inches(w); t.cell(0,i).text=h; shade(t.cell(0,i),NAVY)
        r=t.cell(0,i).paragraphs[0].runs[0]; r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    repeat(t.rows[0])
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=v; cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cells[i])
    return t
def bullet(text): doc.add_paragraph(text,style='List Bullet')
def picture(name,alt):
    doc.add_picture(str(IMG/name),width=Inches(6.45)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    shape=doc.inline_shapes[-1]; shape._inline.docPr.set('descr',alt); shape._inline.docPr.set('title',alt)

doc=Document(); s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11); s.top_margin=Inches(.75); s.bottom_margin=Inches(.7); s.left_margin=Inches(.85); s.right_margin=Inches(.85)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[n]; st.font.name='Arial'; st._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial')
doc.styles['Normal'].font.size=Pt(10.5); doc.styles['Normal'].font.color.rgb=RGBColor.from_string(INK); doc.styles['Normal'].paragraph_format.space_after=Pt(5); doc.styles['Normal'].paragraph_format.line_spacing=1.06
for n,size,color,before,after in [('Title',25,NAVY,0,5),('Subtitle',12,MUTED,0,12),('Heading 1',18,NAVY,15,6),('Heading 2',13,BLUE,10,5)]:
    st=doc.styles[n]; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=n!='Subtitle'; st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
h=s.header.paragraphs[0]; h.text='EXPLORAX / CAZA ESTELAR'; h.runs[0].font.size=Pt(8); h.runs[0].font.bold=True; h.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=f.add_run('Avance del proyecto | '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(MUTED); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(30); r=p.add_run('AVANCE DEL PROYECTO'); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string(BLUE)
doc.add_paragraph('Caza Estelar',style='Title')
doc.add_paragraph('Versión final planeada',style='Subtitle')
table(['Dato','Información'],[('Estudiante','____________________________'),('Curso / sección','____________________________'),('Docente','____________________________'),('Institución','____________________________'),('Fecha','24 de agosto de 2026')],[1.7,4.7])

doc.add_paragraph('1. Aplicación del juego',style='Heading 1')
doc.add_paragraph('1.1 Diagramas de flujo y UML',style='Heading 2')
picture('flujo.png','Flujo planeado de la aplicación Caza Estelar.')
doc.add_paragraph('El juego cargará sus recursos, autenticará al alumno, mostrará el menú y el mapa de misiones, ejecutará la actividad, calculará el resultado y guardará el progreso en la base de datos por medio de la API.')

doc.add_paragraph('1.2 UML de entidades',style='Heading 2')
picture('entidades.png','Entidades principales planeadas para el juego y la administración.')
doc.add_paragraph('Las entidades principales serán Usuario, Alumno, Institución, Catálogo, Misión, Intento y Progreso. Usuario representa la cuenta de acceso común para alumnos, docentes, directores y administradores; Alumno representa exclusivamente el perfil académico vinculado a una cuenta de tipo alumno, con fecha de nacimiento, grado, sección e institución. Esta separación evita guardar datos académicos en cuentas administrativas. En la base de datos también se incluirán Rol, Sección, AsignaciónDocente, Constelación, Actividad, Respuesta, RecursoMultimedia y Auditoría.')

doc.add_paragraph('1.3 UML de variables',style='Heading 2')
picture('variables.png','Variables principales de sesión, configuración, actividad y progreso.')
table(['Clase','Variables planeadas'],[
('PlayerSession','userId, token, authenticated, role, institutionId'),('AppSettings','language, soundEnabled, musicVolume, effectsVolume'),('StarHuntSession','targets, captured, score, elapsedTime, completed'),('MissionProgress','missionId, score, percentage, stars, attempts, completed, updatedAt')],[2.0,4.4])

doc.add_paragraph('1.4 UML de escenarios',style='Heading 2')
picture('escenarios.png','Casos de uso planeados para alumnos y administradores.')
table(['Actor','Escenarios'],[
('Alumno','Iniciar sesión; elegir constelación y actividad; jugar; pausar por desconexión; reanudar; consultar progreso.'),('Docente','Consultar secciones asignadas, resultados, alumnos con actividades incompletas y reportes.'),('Director','Consultar todos los grados y secciones de su institución.'),('Administrador','Gestionar usuarios, instituciones, contenido, bloqueos, publicaciones y reportes.'),('Revisor académico','Revisar y aprobar contenidos antes de publicarlos.')],[1.65,4.75])

doc.add_paragraph('1.5 Programación de eventos',style='Heading 2')
table(['Evento','Respuesta planeada'],[
('Inicio de la aplicación','Precargar recursos, verificar conexión y comprobar una sesión guardada.'),('Inicio de sesión','Enviar credenciales a la API y crear la sesión según el rol.'),('Selección de actividad','Consultar desbloqueos, cargar contenido y comenzar el temporizador.'),('Respuesta del alumno','Evaluar, actualizar puntaje y mostrar retroalimentación inmediata.'),('Pérdida de conexión','Pausar actividad y temporizador; mostrar aviso hasta recuperar la conexión.'),('Fin de actividad','Calcular porcentaje, estrellas y número de intento; enviar el resultado a la API.'),('Cambio de idioma o sonido','Guardar la preferencia y actualizar la interfaz.'),('Bloqueo remoto','Cerrar la sesión de inmediato y eliminar intentos incompletos locales.')],[2.1,4.3])

doc.add_paragraph('1.6 Áreas visuales y distribución de elementos',style='Heading 2')
table(['Pantalla','Distribución planeada'],[
('Carga','Logo y animación al centro; progreso en la parte inferior; fondo espacial.'),('Acceso','Escena de constelaciones al fondo; formulario en panel lateral; ajustes en esquina superior izquierda.'),('Menú','Observatorio al fondo; botones principales agrupados a la derecha; perfil y ajustes visibles.'),('Mapa','Constelaciones en el área central; progreso y bloqueos junto a cada tema.'),('Actividad','Interacción al centro; instrucción arriba; puntaje, tiempo y pausa en los bordes.'),('Resultados','Porcentaje y estrellas al centro; botones para repetir, continuar o volver al mapa.')],[1.35,5.05])

doc.add_paragraph('1.7 Elementos dinámicos',style='Heading 2')
for x in ['Transiciones entre pantallas y viaje estelar.','Estrellas pulsantes, partículas, lluvia estelar y efectos de iluminación.','Mira de telescopio, objetivos, respuestas seleccionables y retroalimentación.','Temporizador, puntaje, porcentaje, estrellas obtenidas y barra de progreso.','Estados de conexión, bloqueo, carga, error, pausa y sincronización.']: bullet(x)
doc.add_paragraph('1.8 Elementos estáticos',style='Heading 2')
for x in ['Fondos del observatorio y las constelaciones.','Logotipos, icono de Android y elementos de marca.','Paneles, marcos, etiquetas, instrucciones y títulos.','Tipografías Fredoka y Nunito.','Botones y recursos visuales en su estado base.']: bullet(x)

doc.add_paragraph('2. Aplicación administrativa',style='Heading 1')
doc.add_paragraph('2.1 Lenguaje de desarrollo',style='Heading 2')
doc.add_paragraph('La aplicación administrativa será web. Se desarrollará con HTML5 y CSS3 para la estructura y presentación, y React con TypeScript para la interfaz y su lógica. El backend se desarrollará con Node.js y NestJS. Esta combinación permite validación de tipos, organización modular y mantenimiento a largo plazo.')

doc.add_paragraph('2.2 Herramientas de terceros',style='Heading 2')
table(['Herramienta','Uso planeado'],[
('React + TypeScript','Interfaz administrativa web.'),('NestJS + Node.js','API REST, reglas de negocio, autenticación y permisos.'),('PostgreSQL','Base de datos central del juego y la administración.'),('Prisma ORM','Acceso tipado y migraciones de la base de datos.'),('Docker Compose','Instalación reproducible de web, API y base de datos en el VPS.'),('Nginx','Proxy inverso, publicación del frontend y terminación HTTPS.'),('Let’s Encrypt','Certificados HTTPS.'),('JWT + Argon2','Sesiones seguras y almacenamiento protegido de contraseñas.'),('ExcelJS y PDFKit','Carga/exportación de Excel y generación de PDF.'),('GitHub Actions','Pruebas y despliegue automatizado al VPS.')],[2.15,4.25])

doc.add_paragraph('2.3 Roles de usuarios',style='Heading 2')
table(['Rol','Permisos planeados'],[
('Alumno','Acceder únicamente al juego y consultar su propio progreso.'),('Docente','Consultar las secciones asignadas y generar reportes académicos.'),('Director','Consultar todos los grados y secciones de su institución.'),('Administrador institucional','Administrar alumnos, docentes y secciones de su institución.'),('Autor de contenido','Crear y editar constelaciones, actividades y recursos.'),('Revisor académico','Aprobar, rechazar y solicitar cambios al contenido.'),('Administrador general','Administrar países, instituciones, usuarios, roles, bloqueos, contenido y auditoría global.')],[1.8,4.6])

doc.add_paragraph('2.4 Utilización de APIs',style='Heading 2')
table(['API / módulo','Funciones planeadas'],[
('Autenticación','Inicio y cierre de sesión, renovación de token, bloqueo y cierre remoto.'),('Usuarios e instituciones','Altas individuales y masivas, grados, secciones y asignaciones.'),('Contenido','Catálogo, constelaciones, actividades, versiones, revisión y publicación.'),('Multimedia','Carga, consulta y descarga de imágenes, audio y video.'),('Progreso','Registro idempotente de intentos, porcentaje, tiempo y estrellas.'),('Reportes','Filtros por fecha, grado, sección y actividad; exportación PDF/Excel.'),('Auditoría','Registro de accesos, bloqueos, cambios y publicaciones.')],[2.0,4.4])

doc.add_paragraph('2.5 Base de datos para el juego',style='Heading 2')
doc.add_paragraph('Se utilizará PostgreSQL como base de datos central porque maneja relaciones, integridad, transacciones, reportes y crecimiento institucional. El juego no se conectará directamente a PostgreSQL: toda operación pasará por la API REST.')
table(['Grupo','Tablas principales'],[
('Identidad','users, roles, permissions, user_roles, sessions'),('Instituciones','countries, institutions, grades, sections, enrollments, teacher_assignments'),('Contenido','catalogs, constellations, activities, questions, answers, media_assets, content_versions'),('Resultados','attempts, activity_results, mission_progress, constellation_progress'),('Operación','sync_events, account_blocks, audit_logs, report_jobs')],[1.55,4.85])
doc.add_paragraph('Relaciones esenciales: una institución posee secciones; una sección agrupa alumnos; una constelación contiene actividades; cada alumno registra múltiples intentos; el progreso conserva el mejor resultado, mientras que attempts mantiene el historial completo.')

doc.add_paragraph('2.6 Lugar de instalación',style='Heading 2')
doc.add_paragraph('La aplicación administrativa, la API y la base de datos se instalarán en un VPS con Ubuntu Server. Los usuarios administrativos ingresarán mediante un dominio HTTPS desde un navegador; no será necesario instalar el sistema en cada computadora.')
doc.add_paragraph('El dominio definido será explorax.app. La solución deberá dimensionarse para una población aproximada de 20,000 usuarios registrados; la capacidad inicial del VPS se ajustará cuando se determine cuántos usuarios podrán conectarse simultáneamente.')
table(['Componente','Instalación en el VPS'],[
('Nginx','Publicará el frontend y dirigirá /api hacia NestJS.'),('Frontend','Archivos HTML, CSS y JavaScript compilados desde React + TypeScript.'),('API','Contenedor Docker con Node.js y NestJS.'),('Base de datos','Contenedor o servicio PostgreSQL con volumen persistente y acceso interno.'),('Archivos','Volumen protegido; posteriormente podrá migrarse a almacenamiento de objetos.'),('Seguridad','Firewall, SSH con llaves, HTTPS, variables secretas y actualizaciones periódicas.'),('Respaldo','Copia diaria cifrada de PostgreSQL y archivos hacia una ubicación externa al VPS.')],[1.65,4.75])

doc.add_paragraph('3. Aspectos que todavía deben aclararse',style='Heading 1')
for x in ['Proveedor, país y capacidad inicial del VPS.','Cantidad máxima estimada de usuarios conectados simultáneamente y volumen de archivos multimedia.','Política de conservación del historial, respaldos y datos personales de menores.','Matriz definitiva de permisos por rol.','Si autor y revisor académico serán personas diferentes obligatoriamente.','Los temas matemáticos, constelaciones y actividades se plantearán posteriormente junto con el equipo académico.','Formato exacto de los archivos Excel/CSV para carga masiva.','Diseño final de los reportes PDF y Excel.','Procedimiento de publicación del APK: descarga privada, MDM institucional o Google Play.']:
    bullet(x)

for t in doc.tables:
    repeat(t.rows[0])
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1
                for r in p.runs:
                    r.font.name='Arial'; r._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial'); r.font.size=Pt(9)
doc.core_properties.title='Avance solicitado - Caza Estelar'
doc.core_properties.subject='Versión final planeada del juego y aplicación administrativa'
doc.core_properties.author='Equipo Caza Estelar'
doc.save(OUT); print(OUT)
